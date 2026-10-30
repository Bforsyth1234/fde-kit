"""Generate the wrangle fixtures (PR 5, §10) from one small, hand-written world.

    python tests/fixtures/wrangle/make_fixtures.py              # CSV, XLSX, expected/*.jsonl
    python tests/fixtures/wrangle/make_fixtures.py --formulas   # also rebuild formulas.xlsx

The world below (plants, parts, months, the numbers) is the ground truth. Each source file is
written FROM the world, with its own messes injected, and `expected/*.jsonl` is written FROM the
world too, never by parsing the files back. That makes the expected facts an independent oracle
for `wrangle.apply`: if apply and this script disagree, one of them misread the mess.

Everything is deterministic (no randomness, no clock): running the script twice gives the same
bytes. XLSX files are re-zipped with a fixed timestamp and a fixed `dcterms:modified`, because
openpyxl stamps both with "now".

formulas.xlsx is the exception. openpyxl can't write cached formula values, and a file with
cached values (plus one formula cell without one) is the case worth testing, so that file is
produced once with LibreOffice and committed. Its sha256 is pinned in FORMULAS_SHA256. The exact
procedure (`build_formulas`, run by `--formulas`):

  1. openpyxl writes formulas_src.xlsx: sheet "Calc", formulas in D3:D7, no cached values.
  2. soffice -env:UserInstallation=file://<tmp>/lo-profile --headless \
         --convert-to xlsx --outdir <tmp>/out <tmp>/formulas_src.xlsx
     LibreOffice recalculates on load and saves every formula with its cached <v> value
     (D6 = 1/0 is cached as the error #DIV/0!).
  3. Inside the resulting zip, in xl/worksheets/sheet1.xml, the <v>…</v> of cell D4 is deleted
     (D4 keeps its <f>B4*C4</f>). Every other zip member is copied unchanged, same order, same
     ZipInfo. openpyxl is NOT used to re-save, because a re-save drops every cached value.
  4. The result is written to formulas.xlsx; its sha256 is printed. Paste it into
     FORMULAS_SHA256 below (LibreOffice output differs between versions and runs, which is why
     the file is committed rather than regenerated in CI).

A normal run only verifies formulas.xlsx against FORMULAS_SHA256 and exits 1 on a mismatch.
"""

from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_FLOOR, ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
EXPECTED = HERE / "expected"

FORMULAS_SHA256 = "e2cdac89742e52fee8c018dde2e702eea6700b2c8d18f8ad1e9520e28fda3f36"

FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
FIXED_W3CDTF = "2026-01-01T00:00:00Z"

# --------------------------------------------------------------------------------------------
# The world
# --------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Plant:
    canonical: str
    erp_code: str  # alias style 1: ERP plant code
    legacy_code: str  # alias style 2: legacy report code
    planner_name: str  # alias style 3: planner's site name
    base_units: int  # typical monthly shipments
    price_eur: Decimal  # average revenue per unit, for the legacy report


PLANTS: tuple[Plant, ...] = (
    Plant("Detroit Stamping 07", "1007", "DET-07", "Detroit Plt 07", 48000, Decimal("23.47")),
    Plant("Toledo Assembly 12", "1012", "TOL-12", "Toledo Plt 12", 36000, Decimal("31.10")),
    Plant("Saltillo Castings 21", "1021", "SAL-21", "Saltillo Plt 21", 22000, Decimal("18.05")),
    Plant("Puebla Machining 33", "1033", "PUE-33", "Puebla Plt 33", 61000, Decimal("12.80")),
    Plant("Monterrey Paint 45", "1045", "MTY-45", "Monterrey Plt 45", 15000, Decimal("9.99")),
)
DET, TOL, SAL, PUE, MTY = PLANTS

PARTS = ("P-1234", "P-2210", "P-3307")

# Which parts each plant has a forecast for (planner_forecast.xlsx).
PLANT_PARTS: dict[str, tuple[str, ...]] = {
    DET.canonical: ("P-1234", "P-2210"),
    TOL.canonical: ("P-1234", "P-3307"),
    SAL.canonical: ("P-2210",),
    PUE.canonical: ("P-1234", "P-2210", "P-3307"),
    MTY.canonical: ("P-3307",),
}

UNKNOWN_LEGACY_PLANT = "GRR-03"  # not in the crosswalk: one row-level reject

Month = tuple[int, int]


def months(start: Month, n: int) -> list[Month]:
    y, m = start
    out = []
    for _ in range(n):
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def month_end(ym: Month) -> date:
    return date(ym[0], ym[1], calendar.monthrange(*ym)[1])


def month_start(ym: Month) -> date:
    return date(ym[0], ym[1], 1)


ERP_MONTHS = months((2025, 10), 12)  # Oct 2025 .. Sep 2026
LEGACY_MONTHS = months((2025, 4), 6)  # Apr 2025 .. Sep 2025 (history before the ERP)
OVERLAP_MONTH: Month = (2025, 10)  # legacy also reports this month for DET (agree) and TOL (not)
PLANNER_MONTHS = months((2026, 11), 4)  # Nov-26 .. Feb-27
FORMULA_PERIOD: Month = (2026, 3)


def erp_units(pi: int, mi: int) -> int:
    """True shipped units for plant index pi in ERP month index mi."""
    return PLANTS[pi].base_units + (pi * 104729 + mi * 7919 + 613) % 2900


def legacy_shipped_k(pi: int, mi: int) -> Decimal:
    """Shipped units in thousands, one decimal, for a legacy-only month."""
    units = PLANTS[pi].base_units - 1500 + (pi * 7919 + mi * 1361) % 2500
    return (Decimal(units) / 1000).quantize(Decimal("0.1"), ROUND_FLOOR)


def legacy_returns(pi: int, mi: int) -> int:
    if (pi, mi) == (3, 2):  # Puebla, Jun 2025: one big return, shown as "(1.204)"
        return -1204
    return -((pi * 13 + mi * 29) % 50 + 1) if (pi + mi) % 3 == 0 else 0


def legacy_revenue(shipped_k: Decimal, returns: int, price: Decimal) -> Decimal:
    return ((shipped_k * 1000 + returns) * price).quantize(Decimal("0.01"), ROUND_HALF_UP)


QUARTERS = (Decimal("0"), Decimal("0.25"), Decimal("0.5"), Decimal("0.75"))


def planner_k(pi: int, part: str, mi: int) -> Decimal | None:
    """Forecast in thousands of units. Quarters only, so every value is exact in binary."""
    if (pi, part, mi) == (4, "P-3307", 3):  # Monterrey P-3307 Feb-27: left blank
        return None
    if (pi, part, mi) == (3, "P-1234", 1):  # Puebla P-1234 Dec-26: big, stored as text
        return Decimal("1012.5")
    qi = PARTS.index(part)
    return Decimal((pi * 37 + qi * 11 + mi * 5) % 60 + 8) + QUARTERS[(pi + mi + qi) % 4]


# formulas.xlsx: (item, qty, unit price); Total = qty * price as a formula.
FORMULA_ITEMS = (("Bolts", 120, 2), ("Nuts", 300, 1), ("Washers", 250, 4), ("Gaskets", 10, 0))
FORMULA_NO_VALUE_CELL = "D4"  # Nuts: cached value stripped
FORMULA_ERROR_CELL = "D6"  # Gaskets: =B6/C6 -> #DIV/0!

# --------------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------------


def json_num(d: Decimal) -> int | float:
    """§6: Decimal -> int if integral, else float."""
    return int(d) if d == d.to_integral_value() else float(d)


def cell_num(d: Decimal) -> int | float:
    """What to store in an XLSX numeric cell."""
    return json_num(d)


def fmt_eu(d: Decimal, places: int) -> str:
    """1234567.891 -> '1.234.567,89' (thousands '.', decimal ',')."""
    return f"{d:,.{places}f}".translate(str.maketrans(",.", ".,"))


def fact(entity, field, value, unit, ym: Month, precision, ref) -> dict:
    return {
        "entity": entity,
        "field": field,
        "value": value,
        "unit": unit,
        "period_start": month_start(ym).isoformat(),
        "period_end": month_end(ym).isoformat(),
        "precision": precision,
        "source_ref": ref,
    }


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8"
    )


def normalise_xlsx(data: bytes) -> bytes:
    """Re-zip with fixed timestamps and a fixed dcterms:created/modified, for stable bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = body.decode("utf-8")
                for tag in ("created", "modified"):
                    text = re.sub(
                        rf"(<dcterms:{tag}[^>]*>)[^<]*(</dcterms:{tag}>)",
                        rf"\g<1>{FIXED_W3CDTF}\g<2>",
                        text,
                    )
                body = text.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, body)
    return out.getvalue()


def save_workbook(wb: openpyxl.Workbook, path: Path) -> None:
    wb.properties.created = datetime(2026, 1, 1)
    wb.properties.creator = "make_fixtures.py"
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(normalise_xlsx(buf.getvalue()))


def write_csv(path: Path, rows: list[list[str]], *, delimiter: str, bom: bool) -> None:
    buf = io.StringIO()
    csv.writer(buf, delimiter=delimiter, lineterminator="\n").writerows(rows)
    path.write_bytes((b"\xef\xbb\xbf" if bom else b"") + buf.getvalue().encode("utf-8"))


# --------------------------------------------------------------------------------------------
# erp_shipments.csv — long, one row per plant-month
# --------------------------------------------------------------------------------------------

ERP_BAD = (1, ERP_MONTHS.index((2026, 3)))  # Toledo Mar 2026: "3?O??" typo -> cell reject
ERP_DUP = (2, ERP_MONTHS.index((2026, 2)))  # Saltillo Feb 2026: re-sent at the end of the file


def build_erp() -> tuple[list[list[str]], list[dict], list[dict]]:
    rows = [["PLANT_CD", "SHIP_DT", "QTY_SHIPPED"]]
    facts: list[dict] = []
    rejects: list[dict] = []
    dup_row: list[str] | None = None
    for mi, ym in enumerate(ERP_MONTHS):
        for pi, p in enumerate(PLANTS):
            units = erp_units(pi, mi)
            code = p.erp_code + " " if p is SAL else p.erp_code  # trailing space
            if p is PUE:
                qty = f"{units:,}"  # "61,234": thousands comma, so the CSV writer quotes it
            elif p is MTY:
                qty = f"{units}  "  # trailing spaces
            else:
                qty = str(units)
            if (pi, mi) == ERP_BAD:
                s = str(units)
                qty = s[:2] + "O" + s[3:]  # letter O typed for a digit
            row = [code, month_end(ym).strftime("%Y%m%d"), qty]
            rows.append(row)
            r = len(rows)
            if (pi, mi) == ERP_BAD:
                rejects.append(
                    {
                        "file": "erp_shipments.csv",
                        "sheet": "erp_shipments",
                        "cell": f"C{r}",
                        "level": "cell",
                        "raw": qty,
                        "stage": "measure",
                        "op": "number",
                        "reason": "number: unparseable",
                    }
                )
                continue
            if (pi, mi) == ERP_DUP:
                dup_row = list(row)
            facts.append(
                fact(p.canonical, "shipped_units", units, "units", ym, None, f"erp_shipments!C{r}")
            )
    assert dup_row is not None
    rows.append(dup_row)  # exact duplicate: duplicates_dropped == 1, first kept
    return rows, facts, rejects


# --------------------------------------------------------------------------------------------
# planner_forecast.xlsx — wide, title rows, two-row merged header, subtotals, (000s)
# --------------------------------------------------------------------------------------------

PLANNER_SHEET = "Plan"
PLANNER_HEADER_UPPER = 4
PLANNER_HEADER_LOWER = 5
PLANNER_FIRST_MONTH_COL = 3  # C


def planner_site(p: Plant) -> str:
    # Toledo is typed in capitals: matches its crosswalk alias only after norm_key (§7.3 step 2).
    return p.planner_name.upper() if p is TOL else p.planner_name


def planner_part_cell(p: Plant, part: str) -> str:
    if p is DET and part == "P-2210":
        return "p-2210"  # lower case; `case: upper` fixes it
    if p is PUE and part == "P-3307":
        return " P-3307 "  # padded; `strip` fixes it
    return part


def month_label(ym: Month) -> str:
    return date(ym[0], ym[1], 1).strftime("%b-%y")  # "Nov-26"


def build_planner(path: Path) -> list[dict]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = PLANNER_SHEET
    ws["A1"] = "Acme Components - S&OP Demand Plan"
    ws["A2"] = "Units (000s)"
    ws["A3"] = "Exported 2026-10-28 from S&OP cycle 14"
    first, last = PLANNER_FIRST_MONTH_COL, PLANNER_FIRST_MONTH_COL + len(PLANNER_MONTHS) - 1
    total_col = last + 1
    ws.cell(PLANNER_HEADER_UPPER, first, "Forecast")
    ws.merge_cells(
        start_row=PLANNER_HEADER_UPPER,
        start_column=first,
        end_row=PLANNER_HEADER_UPPER,
        end_column=last,
    )
    ws.cell(PLANNER_HEADER_LOWER, 1, "Site")
    ws.cell(PLANNER_HEADER_LOWER, 2, "Part No")
    for mi, ym in enumerate(PLANNER_MONTHS):
        ws.cell(PLANNER_HEADER_LOWER, first + mi, month_label(ym))
    ws.cell(PLANNER_HEADER_LOWER, total_col, "Total")

    facts: list[dict] = []
    r = PLANNER_HEADER_LOWER
    grand = [Decimal(0)] * len(PLANNER_MONTHS)
    for pi, p in enumerate(PLANTS):
        sub = [Decimal(0)] * len(PLANNER_MONTHS)
        for part in PLANT_PARTS[p.canonical]:
            r += 1
            ws.cell(r, 1, planner_site(p))
            ws.cell(r, 2, planner_part_cell(p, part))
            row_total = Decimal(0)
            for mi, ym in enumerate(PLANNER_MONTHS):
                v = planner_k(pi, part, mi)
                c = first + mi
                if v is None:
                    continue  # empty cell: empty_skipped
                if v >= 1000:
                    ws.cell(r, c, f"{v:,}")  # number stored as text, "1,012.5"
                else:
                    ws.cell(r, c, cell_num(v))
                sub[mi] += v
                row_total += v
                ref = f"{PLANNER_SHEET}!{get_column_letter(c)}{r}"
                facts.append(
                    fact(
                        f"{p.canonical}/{part}",
                        "demand_forecast_units",
                        json_num(v * 1000),
                        "units",
                        ym,
                        1000,
                        ref,
                    )
                )
            ws.cell(r, total_col, cell_num(row_total))
        r += 1
        ws.cell(r, 1, "Subtotal")
        for mi in range(len(PLANNER_MONTHS)):
            ws.cell(r, first + mi, cell_num(sub[mi]))
            grand[mi] += sub[mi]
        ws.cell(r, total_col, cell_num(sum(sub, Decimal(0))))
    r += 1
    ws.cell(r, 1, "Total")
    for mi in range(len(PLANNER_MONTHS)):
        ws.cell(r, first + mi, cell_num(grand[mi]))
    ws.cell(r, total_col, cell_num(sum(grand, Decimal(0))))

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Forecast is unconstrained demand. Subtotals are per site."
    save_workbook(wb, path)
    return facts


# --------------------------------------------------------------------------------------------
# legacy_report.csv — ';' delimited, decimal ',', DD/MM/YYYY, (parens) returns, (000s) shipped
# --------------------------------------------------------------------------------------------

LEGACY_HEADER = ["Plant", "Report Date", "Shipped (000s)", "Returns", "Revenue (EUR)"]
LEGACY_UNKNOWN_AT = (LEGACY_MONTHS.index((2025, 7)), 2)  # after the 3rd plant of Jul 2025


def legacy_returns_str(n: int) -> str:
    return f"({fmt_eu(Decimal(-n), 0)})" if n < 0 else fmt_eu(Decimal(n), 0)


def build_legacy() -> tuple[list[list[str]], list[dict], list[dict]]:
    rows = [LEGACY_HEADER]
    facts: list[dict] = []
    rejects: list[dict] = []

    def emit(p: Plant | None, code: str, ym: Month, k: Decimal, ret: int, rev: Decimal) -> None:
        rows.append(
            [
                code,
                month_end(ym).strftime("%d/%m/%Y"),
                fmt_eu(k, 1),
                legacy_returns_str(ret),
                fmt_eu(rev, 2),
            ]
        )
        r = len(rows)
        if p is None:
            rejects.append(
                {
                    "file": "legacy_report.csv",
                    "sheet": "legacy_report",
                    "row": r,
                    "level": "row",
                    "raw": code,
                    "stage": "entity",
                    "reason": "unknown entity",
                }
            )
            return
        sheet = "legacy_report"
        facts.append(
            fact(
                p.canonical, "shipped_units", json_num(k * 1000), "units", ym, 1000, f"{sheet}!C{r}"
            )
        )
        facts.append(fact(p.canonical, "returned_units", ret, "units", ym, None, f"{sheet}!D{r}"))
        facts.append(
            fact(p.canonical, "revenue_eur", json_num(rev), "EUR", ym, None, f"{sheet}!E{r}")
        )

    for mi, ym in enumerate(LEGACY_MONTHS):
        for pi, p in enumerate(PLANTS):
            k = legacy_shipped_k(pi, mi)
            ret = legacy_returns(pi, mi)
            emit(p, p.legacy_code, ym, k, ret, legacy_revenue(k, ret, p.price_eur))
            if (mi, pi) == LEGACY_UNKNOWN_AT:
                k2 = Decimal("17.3")
                emit(None, UNKNOWN_LEGACY_PLANT, ym, k2, 0, legacy_revenue(k2, 0, Decimal("11")))

    # The overlap with erp_shipments: Oct 2025 for Detroit and Toledo.
    omi = ERP_MONTHS.index(OVERLAP_MONTH)
    for pi, p in ((0, DET), (1, TOL)):
        erp = Decimal(erp_units(pi, omi))
        k = (erp / 1000).quantize(Decimal("0.1"), ROUND_FLOOR)  # agrees at precision 1000
        if p is TOL:
            k += 3  # THE disagreement: ~3,000 units off
        ret = 0
        emit(p, p.legacy_code, OVERLAP_MONTH, k, ret, legacy_revenue(k, ret, p.price_eur))

    _check_overlap()
    return rows, facts, rejects


def _round_k(d: Decimal) -> Decimal:
    return (d / 1000).quantize(Decimal(1), ROUND_HALF_UP)


def _check_overlap() -> None:
    """Assert the agree / disagree pair really is that under `values_agree` at p=1000."""
    omi = ERP_MONTHS.index(OVERLAP_MONTH)
    for pi, want_agree in ((0, True), (1, False)):
        erp = Decimal(erp_units(pi, omi))
        k = (erp / 1000).quantize(Decimal("0.1"), ROUND_FLOOR) + (0 if want_agree else 3)
        agree = _round_k(erp) == _round_k(k * 1000)
        assert agree is want_agree, (PLANTS[pi].canonical, erp, k)
        if want_agree:
            assert erp != k * 1000, "the agreeing pair should differ exactly, agree at precision"


# --------------------------------------------------------------------------------------------
# formulas.xlsx — committed; see the module docstring
# --------------------------------------------------------------------------------------------


def formulas_expected() -> tuple[list[dict], list[dict]]:
    facts, rejects = [], []
    for i, (item, qty, price) in enumerate(FORMULA_ITEMS):
        r = 3 + i
        ref = f"D{r}"
        if ref == FORMULA_NO_VALUE_CELL:
            rejects.append(
                {
                    "file": "formulas.xlsx",
                    "sheet": "Calc",
                    "cell": ref,
                    "level": "cell",
                    "stage": "measure",
                    "reason": "formula has no cached value",
                }
            )
            continue
        if ref == FORMULA_ERROR_CELL:
            rejects.append(
                {
                    "file": "formulas.xlsx",
                    "sheet": "Calc",
                    "cell": ref,
                    "level": "cell",
                    "raw": "#DIV/0!",
                    "stage": "measure",
                    "reason": "error cell",
                }
            )
            continue
        facts.append(
            fact(item, "line_total_usd", qty * price, "USD", FORMULA_PERIOD, None, f"Calc!{ref}")
        )
    return facts, rejects


def build_formulas(out: Path) -> str:
    soffice = shutil.which("soffice")
    if not soffice:
        sys.exit("soffice not found; install LibreOffice to rebuild formulas.xlsx")
    with tempfile.TemporaryDirectory() as tmp_s:
        tmp = Path(tmp_s)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calc"
        ws["A1"] = "Period: " + date(*FORMULA_PERIOD, 1).strftime("%b %Y")  # "Period: Mar 2026"
        for c, label in enumerate(("Item", "Qty", "Unit Price", "Total"), start=1):
            ws.cell(2, c, label)
        for i, (item, qty, price) in enumerate(FORMULA_ITEMS):
            r = 3 + i
            ws.cell(r, 1, item)
            ws.cell(r, 2, qty)
            ws.cell(r, 3, price)
            ws.cell(r, 4, f"=B{r}/C{r}" if f"D{r}" == FORMULA_ERROR_CELL else f"=B{r}*C{r}")
        tr = 3 + len(FORMULA_ITEMS)
        ws.cell(tr, 1, "Total")
        ws.cell(tr, 4, f"=SUM(D3:D{tr - 2})")
        src = tmp / "formulas_src.xlsx"
        wb.save(src)
        outdir = tmp / "out"
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{tmp}/lo-profile",
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(outdir),
                str(src),
            ],
            check=True,
            capture_output=True,
        )
        converted = (outdir / "formulas_src.xlsx").read_bytes()

    zin = zipfile.ZipFile(io.BytesIO(converted))
    buf = io.BytesIO()
    stripped = 0
    with zipfile.ZipFile(buf, "w") as zout:
        for info in zin.infolist():
            body = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = body.decode("utf-8")
                pat = re.compile(rf'(<c r="{FORMULA_NO_VALUE_CELL}"[^>]*>.*?)<v>[^<]*</v>(.*?</c>)')
                text, stripped = pat.subn(r"\1\2", text)
                body = text.encode("utf-8")
            zout.writestr(info, body)
    if stripped != 1:
        sys.exit(f"expected to strip one <v> from {FORMULA_NO_VALUE_CELL}, stripped {stripped}")
    out.write_bytes(buf.getvalue())
    return hashlib.sha256(buf.getvalue()).hexdigest()


# --------------------------------------------------------------------------------------------


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument(
        "--formulas",
        action="store_true",
        help="rebuild formulas.xlsx with LibreOffice and print its sha256",
    )
    args = ap.parse_args()
    EXPECTED.mkdir(exist_ok=True)

    erp_rows, erp_facts, erp_rejects = build_erp()
    write_csv(HERE / "erp_shipments.csv", erp_rows, delimiter=",", bom=False)
    write_jsonl(EXPECTED / "erp_shipments.facts.jsonl", erp_facts)
    write_jsonl(EXPECTED / "erp_shipments.rejects.jsonl", erp_rejects)

    planner_facts = build_planner(HERE / "planner_forecast.xlsx")
    write_jsonl(EXPECTED / "planner_forecast.facts.jsonl", planner_facts)

    legacy_rows, legacy_facts, legacy_rejects = build_legacy()
    write_csv(HERE / "legacy_report.csv", legacy_rows, delimiter=";", bom=True)
    write_jsonl(EXPECTED / "legacy_report.facts.jsonl", legacy_facts)
    write_jsonl(EXPECTED / "legacy_report.rejects.jsonl", legacy_rejects)

    f_facts, f_rejects = formulas_expected()
    write_jsonl(EXPECTED / "formulas.facts.jsonl", f_facts)
    write_jsonl(EXPECTED / "formulas.rejects.jsonl", f_rejects)

    formulas = HERE / "formulas.xlsx"
    if args.formulas:
        sha = build_formulas(formulas)
        print(f"formulas.xlsx rebuilt; set FORMULAS_SHA256 = {sha!r}")
        return 0
    if not formulas.exists():
        print("formulas.xlsx missing; run with --formulas", file=sys.stderr)
        return 1
    sha = hashlib.sha256(formulas.read_bytes()).hexdigest()
    if sha != FORMULAS_SHA256:
        print(f"formulas.xlsx sha256 {sha} != FORMULAS_SHA256 {FORMULAS_SHA256}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
