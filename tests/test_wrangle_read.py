"""Grid reading (§7.1): coordinates, kinds, merged ranges, formulas without cached values, CSV."""

from __future__ import annotations

import datetime as dt
import hashlib

import openpyxl
import pytest

from fde_kit.wrangle.parse import ParseError, run_ops
from fde_kit.wrangle.read import Cell, col_letter, col_number, parse_ref, read_grid


@pytest.mark.parametrize(
    ("n", "letters"),
    [(1, "A"), (7, "G"), (26, "Z"), (27, "AA"), (52, "AZ"), (53, "BA"), (702, "ZZ"), (703, "AAA")],
)
def test_col_letter_round_trip(n, letters):
    assert col_letter(n) == letters
    assert col_number(letters) == n


def test_col_letter_rejects_zero():
    with pytest.raises(ValueError):
        col_letter(0)


@pytest.mark.parametrize(
    ("ref", "want"),
    [("G42", (42, 7)), ("A1", (1, 1)), ("AA10", (10, 27)), ("$B$3", (3, 2)), ("c5", (5, 3))],
)
def test_parse_ref(ref, want):
    assert parse_ref(ref) == want
    row, col = want
    assert Cell(row, col, "", "empty").ref == ref.replace("$", "").upper()


@pytest.mark.parametrize("ref", ["", "42", "G", "G0", "1G", "Plan!G42", "G42:H43"])
def test_parse_ref_rejects(ref):
    with pytest.raises(ValueError):
        parse_ref(ref)


# --------------------------------------------------------------------------------- xlsx


@pytest.fixture
def plan_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws["A1"] = "Planner forecast (000s)"
    ws.merge_cells("B2:D2")
    ws["B2"] = "2026"
    ws["A3"] = "Site"
    ws["B3"] = "Jan-26"
    ws["C3"] = "Feb-26"
    ws["D3"] = "Total"
    ws["A4"] = "DET-07"
    ws["B4"] = 1234
    ws["C4"] = 0.1
    ws["D4"] = "=B4+C4"  # openpyxl never writes a cached value
    ws["A5"] = "GRY-01"
    ws["B5"] = "#N/A"
    ws["C5"] = "#DIV/0!"
    ws["D5"] = -1234.5
    ws["A6"] = dt.datetime(2026, 3, 4)
    ws["B6"] = dt.datetime(2026, 3, 4, 13, 45)
    ws["C6"] = dt.date(2026, 1, 31)
    ws["D6"] = True
    ws["E7"] = 1e20
    ws["F7"] = 1.5e-7
    ws["A7"] = "   "
    other = wb.create_sheet("Notes")
    other["A1"] = "hello"
    path = tmp_path / "forecast.xlsx"
    wb.save(path)
    return path


def test_xlsx_cells_and_kinds(plan_xlsx):
    g = read_grid(plan_xlsx, "Plan")
    assert g.sheet == "Plan"
    assert g.path == str(plan_xlsx)
    assert g.file_sha256 == hashlib.sha256(plan_xlsx.read_bytes()).hexdigest()

    def at(ref):
        c = g.cell(*parse_ref(ref))
        assert c.ref == ref
        return c.raw, c.kind

    assert at("A1") == ("Planner forecast (000s)", "text")
    assert at("A3") == ("Site", "text")
    assert at("B4") == ("1234", "number")
    assert at("C4") == ("0.1", "number")
    assert at("D5") == ("-1234.5", "number")
    assert at("E7") == ("100000000000000000000", "number")
    assert at("F7") == ("0.00000015", "number")
    assert at("B5") == ("#N/A", "error")
    assert at("C5") == ("#DIV/0!", "error")
    assert at("A6") == ("2026-03-04", "date")
    assert at("B6") == ("2026-03-04T13:45:00", "date")
    assert at("C6") == ("2026-01-31", "date")
    assert at("D6") == ("true", "bool")
    assert at("A7")[1] == "empty"
    assert at("A2") == ("", "empty")


def test_formula_without_cached_value_is_never_zero(plan_xlsx):
    g = read_grid(plan_xlsx, "Plan")
    c = g.cell(4, 4)
    assert (c.ref, c.kind, c.raw) == ("D4", "formula_no_value", "")
    with pytest.raises(ParseError):
        run_ops(c.raw, [{"op": "number"}])


def test_merged_ranges_recorded_not_filled(plan_xlsx):
    g = read_grid(plan_xlsx, "Plan")
    assert g.merged == [(2, 2, 2, 4)]
    assert g.cell(2, 2).raw == "2026"
    assert g.cell(2, 3).kind == "empty"
    assert g.cell(2, 4).kind == "empty"


def test_rows_are_rectangular_and_coordinates_match(plan_xlsx):
    g = read_grid(plan_xlsx, "Plan")
    assert len(g.rows) == 7
    assert {len(r) for r in g.rows} == {6}
    for i, row in enumerate(g.rows, start=1):
        for j, c in enumerate(row, start=1):
            assert (c.row, c.col) == (i, j)
    assert g.cell(100, 100) == Cell(100, 100, "", "empty")


def test_xlsx_numbers_and_dates_parse_with_ops(plan_xlsx):
    g = read_grid(plan_xlsx, "Plan")
    assert run_ops(g.cell(4, 2).raw, [{"op": "number"}]) == 1234
    assert str(run_ops(g.cell(4, 3).raw, [{"op": "number"}])) == "0.1"
    assert run_ops(g.cell(7, 5).raw, [{"op": "number"}]) == 10**20
    assert run_ops(g.cell(6, 1).raw, [{"op": "date", "format": "%Y-%m-%d"}]) == dt.date(2026, 3, 4)


def test_default_sheet_is_first_and_other_sheet_readable(plan_xlsx):
    assert read_grid(plan_xlsx).sheet == "Plan"
    notes = read_grid(plan_xlsx, "Notes")
    assert notes.rows == [[Cell(1, 1, "hello", "text")]]
    assert notes.merged == []


def test_missing_sheet_lists_available(plan_xlsx):
    with pytest.raises(ValueError, match=r"no sheet 'Forecast'.*'Plan', 'Notes'"):
        read_grid(plan_xlsx, "Forecast")


def test_cached_formula_value_is_read(tmp_path):
    """A formula whose cached value is present reads as that value (as LibreOffice saves it).

    openpyxl can't write cached values, so patch one into the sheet XML by hand."""
    import re
    import zipfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["A2"] = "=A1*3"
    ws["A3"] = "=1/0"
    src = tmp_path / "src.xlsx"
    wb.save(src)
    dst = tmp_path / "cached.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode()
                xml = re.sub(r"<f>A1\*3</f>(<v\s*/>|<v></v>)?", "<f>A1*3</f><v>6</v>", xml)
                xml = xml.replace('<c r="A3"', '<c r="A3" t="e"')
                xml = re.sub(r"<f>1/0</f>(<v\s*/>|<v></v>)?", "<f>1/0</f><v>#DIV/0!</v>", xml)
                data = xml.encode()
            zout.writestr(item, data)
    g = read_grid(dst)
    assert (g.cell(2, 1).raw, g.cell(2, 1).kind) == ("6", "number")
    assert (g.cell(3, 1).raw, g.cell(3, 1).kind) == ("#DIV/0!", "error")


def test_unsupported_suffix(tmp_path):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF")
    with pytest.raises(ValueError, match="unsupported"):
        read_grid(p)


# ---------------------------------------------------------------------------------- csv


def test_csv_with_bom_and_comma(tmp_path):
    p = tmp_path / "erp_shipments.csv"
    p.write_bytes('Site,Qty,Date\nDET-07,"1,234",04/03/2026\nGRY-01,,\n'.encode("utf-8-sig"))
    g = read_grid(p)
    assert g.sheet == "erp_shipments"
    assert g.merged == []
    assert g.file_sha256 == hashlib.sha256(p.read_bytes()).hexdigest()
    assert g.cell(1, 1) == Cell(1, 1, "Site", "text")  # BOM stripped
    assert g.cell(2, 2) == Cell(2, 2, "1,234", "text")
    assert g.cell(3, 2) == Cell(3, 2, "", "empty")
    assert {c.kind for r in g.rows for c in r} <= {"text", "empty"}
    assert run_ops(g.cell(2, 2).raw, [{"op": "number"}]) == 1234


def test_csv_semicolon_sniffed(tmp_path):
    p = tmp_path / "legacy_report.csv"
    p.write_text(
        "Site;Datum;Menge\nDET-07;04/03/2026;1.234,5\nGRY-01;05/03/2026;(12,0)\n",
        encoding="utf-8",
    )
    g = read_grid(p)
    assert [c.raw for c in g.rows[0]] == ["Site", "Datum", "Menge"]
    assert g.cell(2, 3).raw == "1.234,5"
    ops = [{"op": "number", "thousands": ".", "decimal": ",", "parens_negative": True}]
    assert str(run_ops(g.cell(3, 3).raw, ops)) == "-12.0"


@pytest.mark.parametrize(("delim", "name"), [("\t", "tab"), ("|", "pipe")])
def test_csv_other_delimiters_sniffed(tmp_path, delim, name):
    p = tmp_path / f"{name}.csv"
    p.write_text(delim.join(["a", "b", "c"]) + "\n" + delim.join(["1", "2", "3"]) + "\n")
    g = read_grid(p)
    assert [c.raw for c in g.rows[1]] == ["1", "2", "3"]


def test_csv_explicit_delimiter_and_encoding(tmp_path):
    p = tmp_path / "latin.csv"
    p.write_bytes("Bez;Wert\nMünchen;3\n".encode("latin-1"))
    g = read_grid(p, encoding="latin-1", delimiter=";")
    assert g.cell(2, 1).raw == "München"


def test_csv_single_column_falls_back_to_comma(tmp_path):
    p = tmp_path / "one.csv"
    p.write_text("value\n1\n2\n")
    g = read_grid(p)
    assert [r[0].raw for r in g.rows] == ["value", "1", "2"]


def test_csv_ragged_rows_padded_and_blank_lines_keep_row_numbers(tmp_path):
    p = tmp_path / "ragged.csv"
    p.write_text("a,b,c\n1\n\n4,5,6,7\n\n\n")
    g = read_grid(p)
    assert len(g.rows) == 4  # trailing blank lines trimmed
    assert {len(r) for r in g.rows} == {4}
    assert g.cell(3, 1).kind == "empty"
    assert g.cell(4, 4) == Cell(4, 4, "7", "text")
    assert g.cell(4, 4).ref == "D4"


def test_csv_sheet_must_match_stem(tmp_path):
    p = tmp_path / "data.csv"
    p.write_text("a\n1\n")
    assert read_grid(p, "data").sheet == "data"
    with pytest.raises(ValueError, match="CSV"):
        read_grid(p, "Plan")


def test_csv_wrong_encoding_is_a_clear_error(tmp_path):
    p = tmp_path / "counties.csv"
    p.write_bytes("STNAME,CTYNAME\nNew Mexico,Doña Ana County\n".encode("latin-1"))
    offset = p.read_bytes().index(b"\xf1")
    with pytest.raises(ValueError) as e:
        read_grid(p)
    msg = str(e.value)
    assert str(p) in msg and f"byte offset {offset}" in msg and "0xf1" in msg
    assert "source.encoding" in msg and "latin-1" in msg and "cp1252" in msg
    assert isinstance(e.value.__cause__, UnicodeDecodeError)
    assert read_grid(p, encoding="latin-1").cell(2, 2).raw == "Doña Ana County"


def test_csv_unknown_encoding(tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("a\n1\n")
    with pytest.raises(ValueError, match="unknown encoding 'no-such-codec'"):
        read_grid(p, encoding="no-such-codec")


def test_xlsx_merged_hidden_values_dropped_and_far_empty_cells_ignored(tmp_path):
    """Only a merged range's top-left value is kept (as openpyxl's full model does), and an
    empty but styled cell far out doesn't widen or lengthen the grid."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Title"
    ws["B1"] = "hidden under the merge"
    ws.merge_cells("A1:C1")
    ws["A2"] = 1
    ws["Z500"].number_format = "0.00"  # styled, no value
    ws["AA2"] = "   "  # whitespace only: empty
    path = tmp_path / "m.xlsx"
    wb.save(path)
    g = read_grid(path)
    assert g.merged == [(1, 1, 1, 3)]
    assert g.cell(1, 2) == Cell(1, 2, "", "empty")
    assert len(g.rows) == 2 and {len(r) for r in g.rows} == {1}  # trimmed to the values


def test_xlsx_reads_only_the_requested_sheet(tmp_path, monkeypatch):
    """The full worksheet model is never built (it took 2.8 GB on a 19 MB workbook)."""
    real = openpyxl.load_workbook
    calls = []

    def spy(*args, **kwargs):
        calls.append(kwargs)
        return real(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", spy)
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    path = tmp_path / "s.xlsx"
    wb.save(path)
    read_grid(path)
    assert calls == [{"read_only": True, "data_only": True}]
