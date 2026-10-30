"""Profile (§7.2) and header labels (§7.6): the fixtures' messes, plus synthetic edge cases."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from fde_kit.wrangle.labels import header_fingerprint, header_labels, normalise_label
from fde_kit.wrangle.parse import compile_ops
from fde_kit.wrangle.profile import Profile, profile
from fde_kit.wrangle.read import Cell, Grid, read_grid

FIX = Path(__file__).parent / "fixtures" / "wrangle"


def _grid(rows: list[list[object]], merged: list[tuple[int, int, int, int]] | None = None) -> Grid:
    """A Grid from Python values: str -> text (or empty when blank), int/float -> number."""
    width = max((len(r) for r in rows), default=0)
    out = []
    for i, r in enumerate(rows, start=1):
        cells = []
        for j in range(1, width + 1):
            v = r[j - 1] if j <= len(r) else ""
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append(Cell(i, j, str(v), "number"))
            else:
                s = str(v)
                cells.append(Cell(i, j, s, "text" if s.strip() else "empty"))
        out.append(cells)
    return Grid(
        path="/somewhere/synthetic.csv",
        sheet="synthetic",
        file_sha256="0" * 64,
        rows=out,
        merged=merged or [],
    )


def _csv(tmp_path: Path, text: str, name: str = "t.csv") -> Grid:
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return read_grid(p)


def _col(p: Profile, label: str):
    return next(c for c in p.columns if c.label == label)


@pytest.fixture(scope="module")
def planner() -> Profile:
    return profile(read_grid(FIX / "planner_forecast.xlsx", "Plan"))


@pytest.fixture(scope="module")
def erp() -> Profile:
    return profile(read_grid(FIX / "erp_shipments.csv"))


@pytest.fixture(scope="module")
def legacy() -> Profile:
    return profile(read_grid(FIX / "legacy_report.csv"))


# ------------------------------------------------------------------------------ labels


def test_header_labels_planner_merged_upper_filled_and_empty_dropped():
    g = read_grid(FIX / "planner_forecast.xlsx", "Plan")
    assert header_labels(g, [4, 5]) == [
        "Site",
        "Part No",
        "Forecast / Nov-26",
        "Forecast / Dec-26",
        "Forecast / Jan-27",
        "Forecast / Feb-27",
        "Total",
    ]
    assert header_labels(g, [5]) == [
        "Site",
        "Part No",
        "Nov-26",
        "Dec-26",
        "Jan-27",
        "Feb-27",
        "Total",
    ]


def test_header_labels_synthetic_cases():
    g = _grid(
        [
            ["Region", "", "Actual", "", "", "Spare"],
            ["", "", "", "", "", ""],
            ["Site", "  Part\n No ", "Q1", "Q2", "", ""],
        ],
        # Region merged down A1:A2 (vertical); Actual merged across C1:E2 (2-D).
        merged=[(1, 1, 2, 1), (1, 3, 2, 5)],
    )
    # Rows 2 and 3: the merges start above the header rows, so they fill their topmost header
    # row (row 2) across; whitespace is collapsed; an all-empty column has an empty label.
    assert header_labels(g, [2, 3]) == [
        "Region / Site",
        "Part No",
        "Actual / Q1",
        "Actual / Q2",
        "Actual",
        "",
    ]
    # Vertical merge over rows 1-2 with both as header rows: filled across only, not down.
    assert header_labels(g, [1, 2, 3])[0] == "Region / Site"
    assert header_labels(g, [1, 2, 3])[5] == "Spare"
    # Order of header_rows doesn't matter; rows outside the grid read as empty.
    assert header_labels(g, [3, 1]) == header_labels(g, [1, 3])
    assert header_labels(g, [3, 99]) == header_labels(g, [3])
    assert header_labels(g, []) == [""] * 6


def test_normalise_label_and_fingerprint():
    assert normalise_label("  Part\tNo \n") == "part no"
    assert normalise_label("STRASSE") == normalise_label("Straße")
    fp = header_fingerprint(["Site", "Part No"])
    assert len(fp) == 12 and int(fp, 16) >= 0
    assert fp == header_fingerprint(["  site", "PART   NO"])
    assert fp != header_fingerprint(["Part No", "Site"])
    assert fp != header_fingerprint(["Site", "Part No", ""])
    want = hashlib.sha256(b"site\x1fpart no").hexdigest()[:12]
    assert fp == want


# ----------------------------------------------------------------------------- planner


def test_planner_header_and_title_rows(planner):
    assert planner.header_rows == [4, 5]
    assert planner.header_found
    assert planner.file == "planner_forecast.xlsx" and planner.sheet == "Plan"
    assert [(t.row, [(c.ref, c.value) for c in t.cells]) for t in planner.title_rows] == [
        (1, [("A1", "Acme Components - S&OP Demand Plan")]),
        (2, [("A2", "Units (000s)")]),
        (3, [("A3", "Exported 2026-10-28 from S&OP cycle 14")]),
    ]
    assert (planner.data.start, planner.data.end) == (6, 20)
    assert planner.notes_rows == []
    assert planner.warnings == []


def test_planner_skip_candidates_by_label_and_sum(planner):
    by_row = {s.row: s for s in planner.skip_candidates}
    assert sorted(by_row) == [8, 11, 13, 17, 19, 20]
    assert all("label" in s.reasons for s in by_row.values())
    # Subtotals over >= 2 rows are also found by sum; row 17 includes D14, the "1,012.5" text.
    for r in (8, 11, 17):
        assert by_row[r].reasons == ["label", "sum"]
        assert "Forecast / Dec-26" in by_row[r].sum_columns
    # One-row windows (Saltillo, Monterrey) are not sum candidates.
    assert by_row[13].reasons == ["label"] and by_row[19].reasons == ["label"]
    # The grand total equals the sum of every non-subtotal row.
    assert by_row[20].reasons == ["label", "sum"]
    assert by_row[20].first_value == "Total"


def test_planner_unpivot_candidate(planner):
    assert len(planner.unpivot_candidates) == 1
    u = planner.unpivot_candidates[0]
    assert (u.from_, u.to, u.format, u.pattern, u.grain) == (
        "Forecast / Nov-26",
        "Forecast / Feb-27",
        "%b-%y",
        None,
        "month",
    )
    assert u.columns == ["C", "D", "E", "F"]
    assert u.ops == [
        {"op": "regex", "pattern": "Forecast / (.+)"},
        {"op": "period", "format": "%b-%y", "grain": "month"},
    ]
    dumped = u.model_dump()
    assert dumped["from"] == "Forecast / Nov-26" and "from_" not in dumped
    run = compile_ops(u.ops)
    assert run("Forecast / Nov-26") == (date(2026, 11, 1), date(2026, 11, 30))
    assert run("Forecast / Feb-27") == (date(2027, 2, 1), date(2027, 2, 28))


def test_planner_columns_and_scale(planner):
    assert [c.label for c in planner.columns][:2] == ["Site", "Part No"]
    site, part = _col(planner, "Site"), _col(planner, "Part No")
    assert site.type == part.type == "text"
    assert site.scale_hint is None and part.scale_hint is None
    # Stats exclude skip rows: Part No is never null among the 9 data rows.
    assert part.null_rate == 0.0 and part.non_null == 9
    assert {"P-1234", "p-2210", " P-3307 "} <= {t.value for t in part.top}
    for c in planner.columns[2:]:
        assert c.type == "number"
        assert c.scale_hint is not None
        assert (c.scale_hint.scale, c.scale_hint.source, c.scale_hint.ref) == (
            1000,
            "title_row",
            "A2",
        )
    assert [(h.scale, h.ref) for h in planner.scale_hints] == [(1000, "A2")]
    dec = _col(planner, "Forecast / Dec-26").number_format
    assert dec is not None
    assert (dec.thousands, dec.decimal, dec.thousands_seen, dec.ambiguous) == (
        ",",
        ".",
        True,
        False,
    )
    feb = _col(planner, "Forecast / Feb-27")
    assert feb.null_rate == round(1 / 9, 4) and feb.non_null == 8
    assert len(feb.samples) == 8 and feb.samples[0].ref == "F6"


def test_planner_keys_and_fingerprint(planner):
    assert planner.keys == [["Site", "Part No"]]
    g = read_grid(FIX / "planner_forecast.xlsx", "Plan")
    assert planner.header_fingerprint == header_fingerprint(header_labels(g, [4, 5]))
    assert profile(g).header_fingerprint == planner.header_fingerprint


# --------------------------------------------------------------------------------- CSVs


def test_erp_profile(erp):
    assert erp.header_rows == [1]
    assert erp.title_rows == [] and erp.skip_candidates == []
    assert (erp.data.start, erp.data.end, erp.data.rows) == (2, 62, 61)
    plant, ship, qty = erp.columns
    assert (plant.label, ship.label, qty.label) == ("PLANT_CD", "SHIP_DT", "QTY_SHIPPED")
    assert ship.type == "date" and ship.date_formats == ["%Y%m%d"]
    assert not ship.day_month_ambiguous
    # Saltillo's code keeps its trailing space; the profile shows the mess verbatim.
    assert {t.value for t in plant.top} == {"1007", "1012", "1021 ", "1033", "1045"}
    assert qty.type == "number"
    nf = qty.number_format
    assert nf is not None
    assert (nf.thousands, nf.decimal, nf.thousands_seen) == (",", ".", True)
    assert nf.parse_rate < 1  # C28 is "38O37"
    assert len(qty.samples) == 10 and qty.samples[0].ref == "C2"
    # The exact duplicate row (24 and 62) means plant+date is not a key.
    assert erp.keys == []
    assert erp.unpivot_candidates == []


def test_legacy_profile(legacy):
    assert legacy.header_rows == [1]
    assert [c.label for c in legacy.columns] == [
        "Plant",
        "Report Date",
        "Shipped (000s)",
        "Returns",
        "Revenue (EUR)",
    ]
    assert _col(legacy, "Plant").type == "text"  # DET-07 is not "DET" currency + -07
    date_col = _col(legacy, "Report Date")
    assert date_col.type == "date" and date_col.date_formats == ["%d/%m/%Y"]
    assert not date_col.day_month_ambiguous

    shipped = _col(legacy, "Shipped (000s)")
    assert shipped.scale_hint is not None
    assert (shipped.scale_hint.scale, shipped.scale_hint.source) == (1000, "label")
    assert _col(legacy, "Returns").scale_hint is None
    assert _col(legacy, "Revenue (EUR)").scale_hint is None

    for label in ("Shipped (000s)", "Returns", "Revenue (EUR)"):
        nf = _col(legacy, label).number_format
        assert nf is not None, label
        assert (nf.thousands, nf.decimal) == (".", ","), label
    returns = _col(legacy, "Returns").number_format
    assert returns.parens_negative and returns.thousands_seen  # "(1.204)"
    revenue = _col(legacy, "Revenue (EUR)").number_format
    assert revenue.thousands_seen and revenue.decimal_seen and not revenue.ambiguous
    assert not _col(legacy, "Shipped (000s)").number_format.parens_negative
    assert legacy.keys == [["Plant", "Report Date"]]


def test_formulas_profile_header_above_total_row():
    p = profile(read_grid(FIX / "formulas.xlsx", "Calc"))
    assert p.header_rows == [2] and p.header_found
    assert [t.cells[0].value for t in p.title_rows] == ["Period: Mar 2026"]
    assert [(s.row, s.reasons) for s in p.skip_candidates] == [(7, ["label"])]
    total = _col(p, "Total")
    assert total.errors == 2 and not total.unique
    assert p.keys == [["Item"]]


# ------------------------------------------------------------------------- determinism


@pytest.mark.parametrize(
    ("name", "sheet"),
    [
        ("planner_forecast.xlsx", "Plan"),
        ("erp_shipments.csv", None),
        ("legacy_report.csv", None),
        ("formulas.xlsx", "Calc"),
    ],
)
def test_profile_json_byte_identical_and_hashed(name, sheet):
    a = profile(read_grid(FIX / name, sheet)).model_dump_json()
    b = profile(read_grid(FIX / name, sheet)).model_dump_json()
    assert a == b
    p = Profile.model_validate_json(a)
    body = p.model_dump_json(exclude={"profile_sha256"})
    assert p.profile_sha256 == hashlib.sha256(body.encode()).hexdigest()
    assert p.model_dump_json() == a  # round-trips, "from" alias included


def test_profile_independent_of_directory(tmp_path):
    src = FIX / "legacy_report.csv"
    copy = tmp_path / "legacy_report.csv"
    copy.write_bytes(src.read_bytes())
    assert profile(read_grid(copy)).model_dump_json() == profile(read_grid(src)).model_dump_json()


def test_profile_sha_changes_with_content(tmp_path):
    a = profile(_csv(tmp_path, "id,qty\nA1,1\nA2,2\n", "a.csv"))
    b = profile(_csv(tmp_path, "id,qty\nA1,1\nA2,3\n", "a.csv"))
    assert a.profile_sha256 != b.profile_sha256
    assert a.header_fingerprint == b.header_fingerprint


# ------------------------------------------------------------------------- synthetic


@pytest.mark.parametrize(
    ("values", "formats", "ambiguous"),
    [
        (["03/04/2026", "05/06/2026", "01/02/2026"], ["%d/%m/%Y", "%m/%d/%Y"], True),
        (["03/04/2026", "13/04/2026", "01/02/2026"], ["%d/%m/%Y"], False),
        (["04/13/2026", "03/04/2026"], ["%m/%d/%Y"], False),
        (["2026-03-04", "2026-12-31"], ["%Y-%m-%d"], False),
    ],
)
def test_day_month_ambiguity(tmp_path, values, formats, ambiguous):
    body = "id,when,qty\n" + "".join(f"r{i},{v},{i + 5}\n" for i, v in enumerate(values))
    col = _col(profile(_csv(tmp_path, body)), "when")
    assert col.type == "date"
    assert col.date_formats == formats
    assert col.day_month_ambiguous is ambiguous


def test_date_formats_tolerate_five_percent():
    rows: list[list[object]] = [["id", "when", "n"]]
    rows += [[f"r{i}", f"{i + 1:02d}/01/2026", i] for i in range(20)]
    rows[5][1] = "n/a"
    col = _col(profile(_grid(rows)), "when")
    assert col.type == "date" and col.date_formats[0] == "%d/%m/%Y"
    rows[6][1] = "unknown"  # 2 of 20 fail: below 95%
    col = _col(profile(_grid(rows)), "when")
    assert col.date_formats == [] and col.type == "text"


def test_subtotal_found_by_sum_without_label():
    g = _grid(
        [
            ["Plant", "Jan", "Feb", "Note"],
            ["A", 1, 2, "x"],
            ["B", 3, 4, "y"],
            ["", 4, 6, ""],  # subtotal with no label
            ["C", 10, 20, "z"],
            ["D", 10, 20, "z"],  # repeats the row before: a one-row window, not a subtotal
            ["E", 5, 5, "w"],
            ["All plants", 29, 51, ""],  # grand total over every non-candidate row
        ]
    )
    p = profile(g)
    assert [(s.row, s.reasons, s.sum_columns) for s in p.skip_candidates] == [
        (4, ["sum"], ["Jan", "Feb"]),
        (8, ["sum"], ["Jan", "Feb"]),
    ]


def test_sum_tolerance_half_unit():
    base: list[list[object]] = [["k", "a", "b"], ["x", 1.2, 2], ["y", 1.2, 2]]
    near = profile(_grid([*base, ["", 2.9, 4.5]]))
    assert [s.row for s in near.skip_candidates] == [4]
    far = profile(_grid([*base, ["", 3.0, 4.6]]))
    assert far.skip_candidates == []


@pytest.mark.parametrize(
    "label", ["Total", "TOTAL plants", "Subtotal", "sub total", "Grand Total", "total:"]
)
def test_total_label_regex(label):
    p = profile(_grid([["k", "v", "w"], ["a", 1, 1], ["b", 2, 2], [label, 99, 98]]))
    assert p.header_found
    hits = [s.row for s in p.skip_candidates if "label" in s.reasons]
    assert hits == ([4] if label != "sub total" else [])


def test_data_region_ends_at_three_empty_rows_and_notes_follow():
    rows: list[list[object]] = [
        ["k", "v", "w"],
        ["a", 1, 1],
        ["", ""],
        ["", ""],
        ["b", 2, 2],  # two empty rows don't end the table
        ["", ""],
        ["", ""],
        ["", ""],
        ["Source: ERP extract", ""],
        ["", "checked by QA"],
    ]
    p = profile(_grid(rows))
    assert p.header_found
    assert (p.data.start, p.data.end, p.data.rows) == (2, 5, 2)
    assert [(n.row, [c.ref for c in n.cells]) for n in p.notes_rows] == [(9, ["A9"]), (10, ["B10"])]
    assert _col(p, "k").null_rate == 0.0  # stats are over non-empty rows


def test_no_header_falls_back_with_warning():
    p = profile(_grid([["", ""], ["alias", "canonical"], ["DET-07", "Detroit"], ["x", "y"]]))
    assert p.header_rows == [2] and not p.header_found
    assert p.warnings and "header rule" in p.warnings[0]


def test_title_rows_above_header_and_upper_needs_merge():
    rows: list[list[object]] = [
        ["Report", "", ""],
        ["", "Forecast", ""],
        ["Site", "Jan", "Feb"],
        ["A", 1, 2],
        ["B", 3, 4],
    ]
    assert profile(_grid(rows)).header_rows == [3]
    assert [t.row for t in profile(_grid(rows)).title_rows] == [1, 2]
    merged = profile(_grid(rows, merged=[(2, 2, 2, 3)]))
    assert merged.header_rows == [2, 3]
    assert [c.label for c in merged.columns] == ["Site", "Forecast / Jan", "Forecast / Feb"]
    assert [t.row for t in merged.title_rows] == [1]


def test_header_window_is_30_rows():
    rows: list[list[object]] = [[f"note {i}", "", ""] for i in range(29)]
    rows += [["k", "v", "w"], ["a", 1, 1], ["b", 2, 2]]
    assert profile(_grid(rows)).header_rows == [30]
    p = profile(_grid([["note", "", ""], *rows]))
    assert not p.header_found and p.header_rows == [1]


def test_two_row_merged_header_from_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Budget FY26"
    ws["B2"] = "Plan"
    ws.merge_cells("B2:E2")
    labels = ["Cost centre", "Q1 FY26", "Q2 FY26", "Q3 FY26", "Q4 FY26"]
    for col, v in zip("ABCDE", labels, strict=True):
        ws[f"{col}3"] = v
    for r, cc in enumerate(["CC-10", "CC-20"], start=4):
        ws[f"A{r}"] = cc
        for j, col in enumerate("BCDE"):
            ws[f"{col}{r}"] = 100 * r + j
    path = tmp_path / "budget.xlsx"
    wb.save(path)
    p = profile(read_grid(path, "S"))
    assert p.header_rows == [2, 3]
    (u,) = p.unpivot_candidates
    assert (u.from_, u.to, u.pattern, u.format, u.grain) == (
        "Plan / Q1 FY26",
        "Plan / Q4 FY26",
        "Q{q} FY{yy}",
        None,
        "quarter",
    )
    run = compile_ops(u.ops)
    assert run("Plan / Q2 FY26") == (date(2026, 4, 1), date(2026, 6, 30))


@pytest.mark.parametrize(
    ("labels", "want"),
    [
        (["k", "2026-01", "2026-02", "2026-03"], ("2026-01", "2026-03", "%Y-%m", "month")),
        (
            ["k", "Jan 2026", "Feb 2026", "Mar 2026", "Total"],
            ("Jan 2026", "Mar 2026", "%b %Y", "month"),
        ),
        (
            ["k", "2026-01-31", "2026-02-28", "2026-03-31"],
            ("2026-01-31", "2026-03-31", "%Y-%m-%d", "month"),
        ),
        (
            ["k", "2026-01-05", "2026-01-12", "2026-01-19"],
            ("2026-01-05", "2026-01-19", "%Y-%m-%d", "day"),
        ),
    ],
)
def test_unpivot_formats_without_prefix(labels, want):
    data = [["x"] + [1] * (len(labels) - 1), ["y"] + [2] * (len(labels) - 1)]
    p = profile(_grid([labels, *data]))
    (u,) = p.unpivot_candidates
    assert (u.from_, u.to, u.format, u.grain) == want
    assert u.ops == [{"op": "period", "format": want[2], "grain": want[3]}]
    compile_ops(u.ops)(u.from_)


def test_unpivot_needs_three_and_one_prefix():
    two = profile(_grid([["k", "Jan-26", "Feb-26", "Total"], ["a", 1, 2, 3], ["b", 1, 2, 3]]))
    assert two.unpivot_candidates == []
    g = _grid(
        [
            ["", "Actual", "", "", "Plan", "", ""],
            ["k", "Jan-26", "Feb-26", "Mar-26", "Jan-26", "Feb-26", "Mar-26"],
            ["a", 1, 2, 3, 4, 5, 6],
            ["b", 1, 2, 3, 4, 5, 6],
        ],
        merged=[(1, 2, 1, 4), (1, 5, 1, 7)],
    )
    p = profile(g)
    assert p.header_rows == [1, 2]
    assert [(u.from_, u.to, u.columns) for u in p.unpivot_candidates] == [
        ("Actual / Jan-26", "Actual / Mar-26", ["B", "C", "D"]),
        ("Plan / Jan-26", "Plan / Mar-26", ["E", "F", "G"]),
    ]
    assert p.unpivot_candidates[1].ops[0] == {"op": "regex", "pattern": "Plan / (.+)"}


def test_unpivot_prefix_regex_is_escaped():
    g = _grid(
        [
            ["", "Units (k)", "", ""],
            ["k", "Jan-26", "Feb-26", "Mar-26"],
            ["a", 1, 2, 3],
            ["b", 1, 2, 3],
        ],
        merged=[(1, 2, 1, 4)],
    )
    (u,) = profile(g).unpivot_candidates
    assert u.ops[0] == {"op": "regex", "pattern": r"Units \(k\) / (.+)"}
    assert compile_ops(u.ops)("Units (k) / Feb-26") == (date(2026, 2, 1), date(2026, 2, 28))


@pytest.mark.parametrize(
    ("text", "scale"),
    [
        ("Units (000s)", 1000),
        ("Sales (000)", 1000),
        ("Revenue '000", 1000),
        ("All figures in thousands", 1000),
        ("Revenue $m", 1_000_000),
        ("Revenue (EUR m)", 1_000_000),
        ("Amounts in millions", 1_000_000),
        ("Length (m)", None),
        ("Units", None),
    ],
)
def test_scale_hint_from_label(text, scale):
    p = profile(_grid([["k", text, "w"], ["a", 1, 1], ["b", 2, 2]]))
    assert p.header_found
    hint = _col(p, text).scale_hint
    assert (hint.scale if hint else None) == scale


def test_scale_hint_from_title_only_on_number_columns():
    p = profile(
        _grid([["All values in millions", "", ""], ["k", "v", "w"], ["a", 1, 2], ["b", 2, 3]])
    )
    assert p.header_rows == [2]
    assert _col(p, "k").scale_hint is None
    hint = _col(p, "v").scale_hint
    assert hint is not None and (hint.scale, hint.source, hint.ref) == (
        1_000_000,
        "title_row",
        "A1",
    )


def test_keys_singles_and_pairs():
    rows: list[list[object]] = [
        ["id", "plant", "month", "qty", "code"],
        ["r1", "A", "Jan", 10, "1007"],
        ["r2", "A", "Feb", 11, "1007"],
        ["r3", "B", "Jan", 12, "1012"],
        ["r4", "B", "Feb", 10, "1012"],
    ]
    p = profile(_grid(rows))
    # id is a key by itself, so no pair with it is listed; qty is a measure, not text-ish;
    # code (a fixed-width digit code) is text-ish and pairs with month.
    assert p.keys == [["id"], ["plant", "month"], ["month", "code"]]
    assert _col(p, "id").unique and not _col(p, "qty").unique
    nulls = profile(_grid([*rows[:2], ["r2", "", "Feb", 1, "1012"], ["r3", "B", "Jan", 2, "1012"]]))
    assert ["plant", "month"] not in nulls.keys


def test_mixed_and_number_style_tie_uses_sheet_preference():
    rows: list[list[object]] = [
        ["k", "eu", "short", "mix"],
        ["a", "1.234,5", "1.234", "x"],
        ["b", "2.000,0", "12", "2"],
    ]
    p = profile(_grid(rows))
    short = _col(p, "short").number_format
    assert short is not None and (short.thousands, short.decimal, short.ambiguous) == (
        ".",
        ",",
        True,
    )
    assert _col(p, "mix").type == "mixed"
    assert _col(p, "mix").number_format is None


def test_currency_values_are_numbers_but_codes_are_not():
    p = profile(
        _grid([["k", "amt", "code"], ["a", "EUR 1.234,50", "DET-07"], ["b", "€ 12,00", "TOL-12"]])
    )
    assert _col(p, "amt").type == "number"
    assert _col(p, "amt").number_format.currency
    assert _col(p, "code").type == "text"
