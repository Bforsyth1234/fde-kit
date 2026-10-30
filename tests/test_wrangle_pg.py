"""Wrangle loads in Postgres (PR 5, §11): the three fixture specs applied to real tenants.

Every spec is copied out of tests/fixtures/wrangle/project into tmp and approved there; the
committed specs stay unapproved. Expected numbers come from `apply.run` (the dry run) in the
test itself, so they are the counts the adapter was handed, not constants copied from the
fixtures README. Each test that writes uses its own tenant (dropped before and after), so a
rerun starts clean.

Needs Postgres with pgvector >= 0.8 at FDE_TEST_PG_URL (an admin/superuser URL), like
tests/test_rag_isolation.py and tests/test_edgar_ingest.py.
"""

from __future__ import annotations

import csv
import io
import os
import re
import shutil
from collections.abc import Callable, Iterator
from datetime import date
from pathlib import Path

import openpyxl
import psycopg
import pytest

from fde_kit.rag import (
    HashEmbedder,
    TenantStore,
    app_url_from_admin,
    create_tenant,
    drop_tenant,
    migrate,
    tenant_search_registry,
)
from fde_kit.rag.tool import reported_line_item
from fde_kit.wrangle.adapter import apply_to_tenant, job_kind_for
from fde_kit.wrangle.apply import ApplyResult, gate, run
from fde_kit.wrangle.spec import LoadedSpec, approve, load_project_spec

PG_URL = os.environ.get("FDE_TEST_PG_URL")
pytestmark = pytest.mark.skipif(not PG_URL, reason="FDE_TEST_PG_URL not set")

FX = Path(__file__).resolve().parent / "fixtures" / "wrangle"
FILES = {
    "erp_shipments": FX / "erp_shipments.csv",
    "planner_forecast": FX / "planner_forecast.xlsx",
    "legacy_report": FX / "legacy_report.csv",
}
SPECS = tuple(FILES)  # load order: ERP first, so the legacy Oct 2025 facts meet it

TENANT_A = "wrangle-pg-a"
TENANT_B = "wrangle-pg-b"
OCT_2025 = date(2025, 10, 1)


@pytest.fixture(scope="module")
def admin_url() -> str:
    assert PG_URL
    with psycopg.connect(PG_URL) as conn:
        try:
            conn.execute("CREATE EXTENSION IF NOT EXISTS vector")
        except psycopg.Error as exc:
            pytest.skip(f"pgvector unavailable: {exc}")
        ver = conn.execute(
            "SELECT extversion FROM pg_extension WHERE extname = 'vector'"
        ).fetchone()
    if ver and tuple(int(x) for x in ver[0].split(".")[:2]) < (0, 8):
        pytest.skip(f"pgvector {ver[0]} < 0.8 (no iterative index scans)")
    migrate(PG_URL)
    return PG_URL


@pytest.fixture(scope="module")
def app_url(admin_url: str) -> str:
    return app_url_from_admin(admin_url)


@pytest.fixture(scope="module")
def store(app_url: str):
    with TenantStore(app_url, HashEmbedder()) as s:
        yield s


@pytest.fixture
def make_tenant(admin_url: str, store: TenantStore) -> Iterator[Callable[[str], str]]:
    """A fresh, empty tenant per call (dropped first in case a previous run left it), dropped
    again when the test ends."""
    made: list[str] = []

    def make(tid: str) -> str:
        drop_tenant(admin_url, tid)
        create_tenant(admin_url, tid, f"Wrangle test {tid}")
        made.append(tid)
        return tid

    yield make
    for tid in made:
        drop_tenant(admin_url, tid)


def _project(base: Path) -> Path:
    """A private copy of the fixture project (specs, dictionary, crosswalks)."""
    dest = base / "project"
    shutil.copytree(FX / "project", dest)
    return dest


def _spec_path(project: Path, name: str) -> Path:
    return project / "specs" / f"{name}.yaml"


def _approved(project: Path, name: str) -> LoadedSpec:
    loaded = approve(_spec_path(project, name), by="test_wrangle_pg")
    assert loaded.approved
    return loaded


def _apply(store: TenantStore, tenant: str, loaded, files: list[Path], **kw):
    with store.session(tenant) as s:
        return apply_to_tenant(s, loaded, files, **kw)


def _counts(store: TenantStore, tenant: str) -> dict[str, int]:
    out: dict[str, int] = {}
    with store.session(tenant) as s:
        for table in ("rag_facts", "rag_documents", "rag_extraction_jobs", "rag_entities"):
            out[table] = s.query(f"SELECT count(*) FROM {table}").rows[0][0]
    return out


def _cells(result: ApplyResult) -> list[str]:
    """ "Sheet!A1" of every fact the dry run produced."""
    return sorted(r.cell for f in result.files for r in f.facts)


def _fetched_cell(content: str, cell: str) -> tuple[str, str]:
    """From fetch_document output: the `row` column and the cell's raw text on the line for
    `cell`'s Excel row. The first line of the output is the `[#1] ...` header."""
    m = re.fullmatch(r"(?:.+!)?([A-Z]+)([0-9]+)", cell)
    assert m, cell
    col, row = m.group(1), m.group(2)
    body = content.split("\n", 1)[1]
    rows = list(csv.reader(io.StringIO(body)))
    header = rows[0]
    assert header[0] == "row"
    line = next(r for r in rows[1:] if r and r[0] == row)
    return line[0], line[header.index(col)]


# ------------------------------------------------------------------ one tenant, three specs


@pytest.fixture(scope="module")
def tenant_a(admin_url: str, store: TenantStore, tmp_path_factory):
    """ERP, planner and legacy applied into TENANT_A, each in its own session (one
    transaction per load, like the CLI). Returns {spec: (ApplyResult, IngestReport)}."""
    drop_tenant(admin_url, TENANT_A)
    create_tenant(admin_url, TENANT_A, "Wrangle test A")
    project = _project(tmp_path_factory.mktemp("wrangle_a"))
    loaded = {}
    for name in SPECS:
        result, report = _apply(store, TENANT_A, _approved(project, name), [FILES[name]])
        assert report is not None, gate(result)
        loaded[name] = (result, report)
    yield loaded
    drop_tenant(admin_url, TENANT_A)


def test_reports_match_the_dry_run(tenant_a):
    for name, (result, report) in tenant_a.items():
        assert report.facts == len(result.facts) > 0, name
        assert (report.documents, report.documents_changed) == (1, 1), name
        assert report.facts_replaced == 0, name
        assert report.job_id, name
    # The fixture README's counts, as a cross-check on the dry run itself.
    assert {n: len(r.facts) for n, (r, _) in tenant_a.items()} == {
        "erp_shipments": 59,
        "planner_forecast": 35,
        "legacy_report": 96,
    }


def test_fact_and_document_counts(tenant_a, store):
    counts = _counts(store, TENANT_A)
    assert counts["rag_facts"] == sum(len(r.facts) for r, _ in tenant_a.values())
    assert counts["rag_documents"] == len(SPECS)
    assert counts["rag_extraction_jobs"] == len(SPECS)
    with store.session(TENANT_A) as s:
        kinds = {r[0] for r in s.query("SELECT kind FROM rag_extraction_jobs").rows}
    assert kinds == {job_kind_for(r.loaded) for r, _ in tenant_a.values()}
    assert all(k.startswith("wrangle:") for k in kinds)


def test_every_stored_fact_cites_its_cell(tenant_a, store):
    """The stored source_refs are exactly the dry run's, cell first (§7.4), so
    `reported_line_item` gives the cell."""
    with store.session(TENANT_A) as s:
        facts = s.facts()
    want = sorted(f.row.source_ref for res, _ in tenant_a.values() for f in res.facts)
    assert sorted(f.source_ref for f in facts) == want
    assert sorted(reported_line_item(f.source_ref) for f in facts) == sorted(
        c for res, _ in tenant_a.values() for c in _cells(res)
    )


@pytest.mark.parametrize(
    ("entity", "field", "period", "cell", "title"),
    [
        (
            "Detroit Stamping 07",
            "shipped_units",
            "2025-11-01",  # Oct 2025 also has a legacy fact; Nov is ERP only
            "erp_shipments!C7",
            "erp_shipments.csv — erp_shipments",
        ),
        (
            "Puebla Machining 33/P-1234",
            "demand_forecast_units",
            "2026-12-01",
            "Plan!D14",
            "planner_forecast.xlsx — Plan",
        ),
        (
            "Toledo Assembly 12",
            "returned_units",
            "2025-10-01",
            "legacy_report!D34",
            "legacy_report.csv — legacy_report",
        ),
    ],
)
def test_query_table_shows_the_cell_as_line_item(
    tenant_a, store, entity, field, period, cell, title
):
    sql = (
        "SELECT f.fact_id, e.canonical_name, f.field, f.period_start, f.value_num "
        "FROM rag_facts f JOIN rag_entities e ON e.entity_id = f.entity_id "
        f"WHERE e.canonical_name = '{entity}' AND f.field = '{field}' "
        f"AND f.period_start = '{period}'"
    )
    out = tenant_search_registry(store, TENANT_A, tools=("query_table",)).execute(
        "query_table", {"sql": sql}
    )
    assert not out.is_error, out.content
    rows = out.content.splitlines()[1:]
    assert len(rows) == 1, out.content
    assert rows[0].startswith("[#1] ")
    assert rows[0].endswith(f"(source: {title}; reported as {cell})"), rows[0]
    assert len(out.sources) == 1 and out.sources[0].tenant_id == TENANT_A


@pytest.mark.parametrize(
    ("entity", "field", "cell", "raw"),
    [
        ("Detroit Stamping 07", "shipped_units", "erp_shipments!C2", "48613"),
        ("Toledo Assembly 12", "shipped_units", "legacy_report!C34", "39,9"),
        ("Puebla Machining 33/P-1234", "demand_forecast_units", "Plan!D14", "1,012.5"),
    ],
)
def test_fetch_document_shows_the_excel_row(tenant_a, store, entity, field, cell, raw):
    """fetch_document on a fact's doc: the line for the fact's row carries the Excel row
    number in the `row` column and the cell's text in the cell's column (§7.4)."""
    with store.session(TENANT_A) as s:
        candidates = s.facts(entity=entity, field=field)
    fact = next(f for f in candidates if reported_line_item(f.source_ref) == cell)
    assert fact.doc_id
    out = tenant_search_registry(store, TENANT_A, tools=("fetch_document",)).execute(
        "fetch_document", {"doc_id": fact.doc_id, "char_start": 0, "char_end": 100_000}
    )
    assert not out.is_error, out.content
    assert out.content.startswith("[#1] ") and " v1 " in out.content.splitlines()[0]
    row = re.search(r"[0-9]+$", cell).group(0)
    assert _fetched_cell(out.content, cell) == (row, raw)


# ------------------------------------------------------------------ conflicts across sources


def _oct_2025(store: TenantStore, entity: str):
    with store.session(TENANT_A) as s:
        facts = s.facts(entity=entity, field="shipped_units")
    return [f for f in facts if f.period_start == OCT_2025]


def test_toledo_erp_and_legacy_form_one_conflict_group(tenant_a, store):
    facts = _oct_2025(store, "Toledo Assembly 12")
    assert sorted(reported_line_item(f.source_ref) for f in facts) == [
        "erp_shipments!C3",
        "legacy_report!C34",
    ]
    groups = {f.conflict_group_id for f in facts}
    assert len(groups) == 1 and None not in groups
    assert sorted(int(f.value) for f in facts) == [36942, 39900]


def test_detroit_agrees_at_the_declared_precision(tenant_a, store):
    facts = _oct_2025(store, "Detroit Stamping 07")
    assert sorted(reported_line_item(f.source_ref) for f in facts) == [
        "erp_shipments!C2",
        "legacy_report!C33",
    ]
    assert [f.conflict_group_id for f in facts] == [None, None]


def test_toledo_is_the_only_conflict_group(tenant_a, store):
    for result, _ in tenant_a.values():
        assert result.conflicts == []  # nothing conflicts within one load
    with store.session(TENANT_A) as s:
        grouped = [f for f in s.facts() if f.conflict_group_id]
    assert len({f.conflict_group_id for f in grouped}) == 1
    assert {f.entity for f in grouped} == {"Toledo Assembly 12"}
    assert len(grouped) == 2


def test_query_table_marks_the_conflict(tenant_a, store):
    out = tenant_search_registry(store, TENANT_A, tools=("query_table",)).execute(
        "query_table",
        {
            "sql": "SELECT f.fact_id, f.value_num FROM rag_facts f "
            "JOIN rag_entities e ON e.entity_id = f.entity_id "
            "WHERE e.canonical_name IN ('Toledo Assembly 12', 'Detroit Stamping 07') "
            "AND f.field = 'shipped_units' AND f.period_start = '2025-10-01' "
            "ORDER BY f.value_num"
        },
    )
    assert not out.is_error, out.content
    rows = out.content.splitlines()[1:]
    assert len(rows) == 4
    by_cell = {re.search(r"reported as ([^)]+)\)", r).group(1): r for r in rows}
    assert by_cell["erp_shipments!C3"].endswith(" conflict=true")
    assert by_cell["legacy_report!C34"].endswith(" conflict=true")
    assert not by_cell["erp_shipments!C2"].endswith(" conflict=true")
    assert not by_cell["legacy_report!C33"].endswith(" conflict=true")


# ------------------------------------------------------------------ re-export


def test_reexport_supersedes_and_replaces_its_facts(make_tenant, store, tmp_path):
    """A new planner export with one cell changed, under the fixed document_key: a second
    version of the one document, its facts replaced, one value different."""
    tid = make_tenant("wrangle-pg-reexport")
    loaded = _approved(_project(tmp_path), "planner_forecast")
    first, r1 = _apply(store, tid, loaded, [FILES["planner_forecast"]])
    assert r1 is not None and r1.facts == len(first.facts) == 35
    with store.session(tid) as s:
        before = {(f.entity, f.field, f.period_start): f.value for f in s.facts()}
        old_doc = {f.doc_id for f in s.facts()}
    assert len(before) == len(first.facts) and len(old_doc) == 1

    export = tmp_path / "planner_forecast.xlsx"  # same name, as a re-export would be
    shutil.copy(FILES["planner_forecast"], export)
    wb = openpyxl.load_workbook(export)  # no formulas in this file, so a resave is safe
    ws = wb["Plan"]
    old_value = ws["C6"].value
    assert isinstance(old_value, int | float)
    ws["C6"] = old_value + 1  # Detroit P-1234 Nov-26, in thousands
    wb.save(export)

    second, r2 = _apply(store, tid, loaded, [export])
    assert r2 is not None, gate(second)
    assert second.documents[0].source == first.documents[0].source == "planner-forecast#Plan"
    assert (r2.documents, r2.documents_changed) == (1, 1)
    assert r2.facts_replaced == len(first.facts)
    assert r2.facts == len(second.facts) == len(first.facts)
    assert r2.job_id and r2.job_id != r1.job_id

    with store.session(tid) as s:
        facts = s.facts()
        after = {(f.entity, f.field, f.period_start): f.value for f in facts}
        new_doc = {f.doc_id for f in facts}
        assert len(new_doc) == 1 and new_doc != old_doc
        new = s.fetch(next(iter(new_doc)), 0, 10)
        old = s.fetch(next(iter(old_doc)), 0, 10)
    assert (new.version, new.superseded_by) == (2, None)
    assert (old.version, old.superseded_by) == (1, new.doc_id)
    assert all(f.conflict_group_id is None for f in facts)
    assert all(reported_line_item(f.source_ref) for f in facts)

    assert _counts(store, tid)["rag_facts"] == len(first.facts)
    assert after.keys() == before.keys()
    diff = {k: (before[k], after[k]) for k in before if before[k] != after[k]}
    key = ("Detroit Stamping 07/P-1234", "demand_forecast_units", date(2026, 11, 1))
    assert list(diff) == [key]
    assert int(diff[key][1]) - int(diff[key][0]) == 1000  # scale 1000


# ------------------------------------------------------------------ idempotence


def test_applying_unchanged_files_twice_writes_nothing(make_tenant, store, tmp_path):
    tid = make_tenant("wrangle-pg-idem")
    project = _project(tmp_path)
    for name in ("erp_shipments", "legacy_report"):
        loaded = _approved(project, name)
        _, first = _apply(store, tid, loaded, [FILES[name]])
        assert first is not None and first.facts > 0 and first.job_id
        before = _counts(store, tid)
        _, again = _apply(store, tid, loaded, [FILES[name]])
        assert again is not None
        assert (again.documents, again.documents_changed) == (1, 0), name
        assert (again.facts, again.facts_replaced, again.job_id) == (0, 0, ""), name
        assert _counts(store, tid) == before, name


# ------------------------------------------------------------------ gates


def test_gates_refuse_and_write_nothing(make_tenant, store, tmp_path):
    tid = make_tenant("wrangle-pg-gates")
    project = _project(tmp_path)
    # Something already loaded, so "unchanged" isn't trivially "empty".
    _, report = _apply(store, tid, _approved(project, "legacy_report"), [FILES["legacy_report"]])
    assert report is not None and report.facts > 0
    baseline = _counts(store, tid)
    erp = _spec_path(project, "erp_shipments")
    erp_file = [FILES["erp_shipments"]]

    # 1. Unapproved: the copied spec has no approval block.
    assert not load_project_spec(erp).approved
    result, report = _apply(store, tid, erp, erp_file)
    assert report is None
    assert any("not approved" in r for r in gate(result))
    assert _counts(store, tid) == baseline

    # 2. Stale: approved, then a crosswalk byte changes, so the bundle hash moves.
    approve(erp, by="test_wrangle_pg")
    with (project / "crosswalks" / "plants.csv").open("a", encoding="utf-8") as fh:
        fh.write("9999,Detroit Stamping 07\n")
    stale = load_project_spec(erp)
    assert stale.spec.approval is not None and not stale.approved
    result, report = _apply(store, tid, erp, erp_file)
    assert report is None
    assert any("stale" in r for r in gate(result))
    assert _counts(store, tid) == baseline

    # 3. Reject rate over the limit: re-approve, then a limit below ERP's 1/60.
    loaded = _approved(project, "erp_shipments")
    result, report = _apply(store, tid, loaded, erp_file, max_reject_rate=0.001)
    assert report is None
    assert result.reject_rate > 0.001
    assert any("reject rate" in r for r in gate(result, max_reject_rate=0.001))
    assert _counts(store, tid) == baseline

    # Control: the same approved spec under the default limit does load.
    _, report = _apply(store, tid, loaded, erp_file)
    assert report is not None and report.facts == len(run(loaded, erp_file).facts)
    assert _counts(store, tid)["rag_facts"] == baseline["rag_facts"] + report.facts


# ------------------------------------------------------------------ isolation


def test_second_tenant_sees_none_of_the_first(tenant_a, make_tenant, store, tmp_path):
    """Tenant B loads only the planner, under the same fixed document_key as A's planner. It
    gets its own version 1, and none of A's facts, documents or cells are visible to it."""
    tid = make_tenant(TENANT_B)
    assert set(_counts(store, tid).values()) == {0}
    loaded = _approved(_project(tmp_path), "planner_forecast")
    result, report = _apply(store, tid, loaded, [FILES["planner_forecast"]])
    assert report is not None and report.facts == len(result.facts)

    with store.session(TENANT_A) as s:
        a_facts = s.facts()
    a_ids = {f.fact_id for f in a_facts}
    a_docs = {f.doc_id for f in a_facts}
    assert len(a_docs) == len(SPECS)

    with store.session(tid) as s:
        b_facts = s.facts()
        assert {f.fact_id for f in b_facts}.isdisjoint(a_ids)
        assert s.facts_by_id(a_ids) == {}
        assert s.facts(field="shipped_units") == []
        for doc_id in a_docs:
            if doc_id in {f.doc_id for f in b_facts}:
                pytest.fail("tenant B shares a doc_id with tenant A")
            with pytest.raises(KeyError):
                s.fetch(doc_id)
    with store.session(tid) as s:
        sources = [r[0] for r in s.query("SELECT source FROM rag_documents").rows]
        versions = [r[0] for r in s.query("SELECT version FROM rag_documents").rows]
    assert sources == ["planner-forecast#Plan"] and versions == [1]
    assert len(b_facts) == len(result.facts)
    assert {reported_line_item(f.source_ref).split("!")[0] for f in b_facts} == {"Plan"}

    reg = tenant_search_registry(store, tid, tools=("query_table", "fetch_document"))
    out = reg.execute(
        "query_table",
        {
            "sql": "SELECT f.fact_id, f.source_ref FROM rag_facts f "
            "WHERE f.source_ref LIKE 'erp_shipments!%' OR f.source_ref LIKE 'legacy_report!%'"
        },
    )
    assert not out.is_error, out.content
    assert out.content.splitlines()[1:] == [] and out.sources == ()
    erp_doc = next(f.doc_id for f in a_facts if f.source_ref.startswith("erp_shipments!"))
    out = reg.execute("fetch_document", {"doc_id": erp_doc})
    assert out.is_error and "No document" in out.content

    # Control: A still sees its own ERP facts through the same statement.
    out = tenant_search_registry(store, TENANT_A, tools=("query_table",)).execute(
        "query_table",
        {"sql": "SELECT f.fact_id FROM rag_facts f WHERE f.source_ref LIKE 'erp_shipments!%'"},
    )
    assert not out.is_error and len(out.content.splitlines()) - 1 == len(
        tenant_a["erp_shipments"][0].facts
    )
