"""Read a CSV or XLSX sheet into a `Grid`: every cell as a string, with its coordinates.

Everything downstream (profile, spec ops, provenance) works on strings with a `kind`, never on
the Python objects openpyxl returns. That keeps one parse path for CSV and XLSX, and it makes
`source_ref` read-back exact: re-reading the cell gives the same string, which the same ops turn
into the same value.

Two XLSX cases matter because getting them wrong silently loads a wrong number:

- A formula cell with no cached value (a file written by a script and never opened in Excel)
  reads as `None` under `data_only=True`. It's `formula_no_value`, so a spec that reads it
  rejects the cell instead of loading 0 or skipping it as empty.
- Excel error values (`#N/A`, `#REF!`, ...) are `error`, never text that a `map` could rescue.

Numbers are `str(Decimal(repr(float)))` in plain notation (no exponent), so `0.1` is `"0.1"`
and not `0.1000000000000000055...`, and the `number` op's strict grammar accepts it.

`merged` ranges are `(min_row, min_col, max_row, max_col)`, 1-based and inclusive. Only the
top-left cell of a range holds a value; `read` doesn't fill the others (header joining does).
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Literal

CellKind = Literal["empty", "text", "number", "date", "bool", "error", "formula_no_value"]

EXCEL_ERRORS = frozenset(
    {
        "#N/A",
        "#REF!",
        "#DIV/0!",
        "#VALUE!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
        "#SPILL!",
        "#CALC!",
        "#GETTING_DATA",
    }
)
XLSX_SUFFIXES = frozenset({".xlsx", ".xlsm"})
CSV_SUFFIXES = frozenset({".csv", ".tsv", ".txt"})
SNIFF_BYTES = 64 * 1024
SNIFF_DELIMITERS = ",;\t|"

_REF_RE = re.compile(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)")


def col_letter(n: int) -> str:
    """1 -> "A", 26 -> "Z", 27 -> "AA"."""
    if n < 1:
        raise ValueError(f"column numbers are 1-based, got {n}")
    letters = ""
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters


def col_number(letters: str) -> int:
    """ "A" -> 1, "AA" -> 27. The inverse of `col_letter`."""
    if not letters.isalpha() or not letters.isascii():
        raise ValueError(f"not a column: {letters!r}")
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n


def parse_ref(ref: str) -> tuple[int, int]:
    """ "G42" -> (42, 7), as (row, col). `$` anchors and lower case are accepted."""
    m = _REF_RE.fullmatch(ref.strip())
    if not m:
        raise ValueError(f"not an A1 cell reference: {ref!r}")
    return int(m.group(2)), col_number(m.group(1))


@dataclass(frozen=True)
class Cell:
    row: int
    col: int
    """1-based, as Excel shows them."""
    raw: str
    """Text as stored; numbers as plain-notation str(Decimal); dates ISO."""
    kind: CellKind

    @property
    def ref(self) -> str:
        return f"{col_letter(self.col)}{self.row}"


@dataclass(frozen=True)
class Grid:
    path: str
    sheet: str
    """The XLSX sheet name; for CSV, the file stem."""
    file_sha256: str
    rows: list[list[Cell]]
    """Rectangular: every row has the same length, padded with empty cells."""
    merged: list[tuple[int, int, int, int]]
    """(min_row, min_col, max_row, max_col), 1-based, inclusive; empty for CSV."""

    def cell(self, row: int, col: int) -> Cell:
        """The cell at 1-based (row, col); an empty cell outside the read area."""
        if 1 <= row <= len(self.rows) and 1 <= col <= len(self.rows[row - 1]):
            return self.rows[row - 1][col - 1]
        return Cell(row, col, "", "empty")


def read_grid(
    path: str | Path,
    sheet: str | None = None,
    *,
    encoding: str = "utf-8-sig",
    delimiter: str | None = None,
) -> Grid:
    """Read one sheet. XLSX: `sheet` defaults to the first one. CSV: `sheet` must be omitted or
    equal the file stem; `encoding` and `delimiter` apply to CSV only."""
    p = Path(path)
    data = p.read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    suffix = p.suffix.lower()
    if suffix in XLSX_SUFFIXES:
        return _read_xlsx(p, sheet, sha)
    if suffix in CSV_SUFFIXES:
        if sheet is not None and sheet != p.stem:
            raise ValueError(f"{p.name} is a CSV; its only sheet is {p.stem!r}, not {sheet!r}")
        return _read_csv(p, data, sha, encoding, delimiter)
    raise ValueError(f"unsupported file type {p.suffix!r} (want .csv or .xlsx): {p}")


def _trim(rows: list[list[Cell]]) -> list[list[Cell]]:
    """Drop trailing rows and columns that are entirely empty, then pad to a rectangle."""
    while rows and all(c.kind == "empty" for c in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for c in reversed(r):
            if c.kind != "empty":
                width = max(width, c.col)
                break
    out = []
    for i, r in enumerate(rows, start=1):
        r = r[:width]
        r.extend(Cell(i, j, "", "empty") for j in range(len(r) + 1, width + 1))
        out.append(r)
    return out


def _plain_decimal(d: Decimal) -> str:
    s = str(d)
    return format(d, "f") if "E" in s or "e" in s else s


def _xlsx_cell(row: int, col: int, value: object, is_formula: bool) -> Cell:
    if value is None:
        return Cell(row, col, "", "formula_no_value" if is_formula else "empty")
    if isinstance(value, bool):
        return Cell(row, col, "true" if value else "false", "bool")
    if isinstance(value, int):
        return Cell(row, col, str(value), "number")
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return Cell(row, col, repr(value), "error")
        return Cell(row, col, _plain_decimal(Decimal(repr(value))), "number")
    if isinstance(value, Decimal):
        return Cell(row, col, _plain_decimal(value), "number")
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return Cell(row, col, value.date().isoformat(), "date")
        return Cell(row, col, value.isoformat(), "date")
    if isinstance(value, dt.date):
        return Cell(row, col, value.isoformat(), "date")
    if isinstance(value, dt.time):
        return Cell(row, col, value.isoformat(), "text")
    text = str(value)
    if text in EXCEL_ERRORS:
        return Cell(row, col, text, "error")
    return Cell(row, col, text, "empty" if not text.strip() else "text")


def _read_xlsx(p: Path, sheet: str | None, sha: str) -> Grid:
    """Stream one sheet's XML once.

    The workbook is opened `read_only=True, data_only=True`, which parses the shared strings and
    styles but no sheet; then only the requested sheet is parsed, by openpyxl's own sheet parser,
    so cell values are exactly what `load_workbook(data_only=True)` gives. The same pass notes the
    cells that hold a formula but no cached value (`<f>` with no or an empty `<v>`) and the
    sheet's `<mergeCells>`, which read-only worksheets don't expose. Memory is the grid itself,
    not openpyxl's worksheet model, and nothing is read twice.
    """
    try:
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
        from openpyxl.worksheet._reader import FORMULA_TAG, WorkSheetParser
    except ImportError as e:  # pragma: no cover - depends on the install
        raise ImportError("reading XLSX needs openpyxl: pip install 'fde-kit[wrangle]'") from e

    class _Parser(WorkSheetParser):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.no_cached_value: set[tuple[int, int]] = set()

        def parse_cell(self, element):
            cell = super().parse_cell(element)
            if cell["value"] is None and element.find(FORMULA_TAG) is not None:
                self.no_cached_value.add((cell["row"], cell["column"]))
            return cell

    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        name = sheet if sheet is not None else names[0]
        if name not in names:
            raise ValueError(f"{p.name} has no sheet {name!r}; available sheets: {names}")
        ws = wb[name]
        if not hasattr(ws, "_get_source"):  # a chartsheet
            raise ValueError(f"{p.name}: sheet {name!r} is not a worksheet")
        sparse: dict[tuple[int, int], Cell] = {}
        with ws._get_source() as src:
            parser = _Parser(
                src,
                ws._shared_strings,
                data_only=True,
                epoch=wb.epoch,
                date_formats=wb._date_formats,
                timedelta_formats=wb._timedelta_formats,
            )
            for _, row in parser.parse():
                for c in row:
                    r_idx, c_idx = c["row"], c["column"]
                    sparse[(r_idx, c_idx)] = _xlsx_cell(
                        r_idx, c_idx, c["value"], (r_idx, c_idx) in parser.no_cached_value
                    )
            merge_refs = parser.merged_cells.mergeCell if parser.merged_cells else []
    finally:
        wb.close()

    merged = []
    for mc in merge_refs:
        min_col, min_row, max_col, max_row = range_boundaries(mc.ref)
        merged.append((min_row, min_col, max_row, max_col))
        # As openpyxl's full model does: only the top-left cell of a merged range keeps a value.
        if (max_row - min_row + 1) * (max_col - min_col + 1) <= len(sparse):
            for r_idx in range(min_row, max_row + 1):
                for c_idx in range(min_col, max_col + 1):
                    if (r_idx, c_idx) != (min_row, min_col):
                        sparse.pop((r_idx, c_idx), None)
        else:
            for key in [k for k in sparse if min_row <= k[0] <= max_row]:
                if min_col <= key[1] <= max_col and key != (min_row, min_col):
                    del sparse[key]
    merged.sort()

    # Size from the non-empty cells only, so styled-but-empty cells far out don't widen the grid.
    height = width = 0
    for (r_idx, c_idx), cell in sparse.items():
        if cell.kind != "empty":
            height = max(height, r_idx)
            width = max(width, c_idx)
    rows = [
        [
            sparse.get((r_idx, c_idx)) or Cell(r_idx, c_idx, "", "empty")
            for c_idx in range(1, width + 1)
        ]
        for r_idx in range(1, height + 1)
    ]
    return Grid(path=str(p), sheet=name, file_sha256=sha, rows=_trim(rows), merged=merged)


def sniff_delimiter(sample: str) -> str:
    """The delimiter csv.Sniffer picks from `,;\\t|`, or "," when it can't tell."""
    try:
        return csv.Sniffer().sniff(sample, delimiters=SNIFF_DELIMITERS).delimiter
    except csv.Error:
        return ","


def _read_csv(p: Path, data: bytes, sha: str, encoding: str, delimiter: str | None) -> Grid:
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError as e:
        bad = data[e.start : e.start + 1].hex()
        raise ValueError(
            f"{p}: not valid {encoding} at byte offset {e.start} (byte 0x{bad}); "
            "set `source.encoding` in the spec to the file's encoding (e.g. latin-1 or cp1252)"
        ) from e
    except LookupError as e:
        raise ValueError(f"{p}: unknown encoding {encoding!r}") from e
    if delimiter is None:
        delimiter = sniff_delimiter(text[:SNIFF_BYTES])
    rows: list[list[Cell]] = []
    for r_idx, record in enumerate(
        csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), start=1
    ):
        rows.append(
            [
                Cell(r_idx, c_idx, v, "empty" if not v.strip() else "text")
                for c_idx, v in enumerate(record, start=1)
            ]
        )
    return Grid(path=str(p), sheet=p.stem, file_sha256=sha, rows=_trim(rows), merged=[])
