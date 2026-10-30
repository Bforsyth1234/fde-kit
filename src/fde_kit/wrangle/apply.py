"""Apply a mapping spec to files: cells -> facts with cell-level provenance (§7.4).

`run(spec, files)` is pure: no database. It is the dry run, and the adapter (`adapter.py`)
loads what it returns, so a dry run and an apply can never disagree about a value.

A load's units are its files, or with `source.sheets` each selected sheet of each XLSX file,
in workbook order (`unit_sheets`). Each unit is read into a `Grid`, then:

1. `header_fingerprint`, if the spec sets one, must match (`force=True` skips only this check).
2. Column labels are built from `header_rows` (`labels.header_labels`); a column the spec names
   that isn't there is a `SpecMismatch`, and the whole file is refused.
3. The data rows run from `first_data_row`, narrowed by `start_after` / `stop_before` when the
   spec sets them (a bound that matches no row is a `SpecMismatch`). Each data row: `skip_rows`
   (counted, not rejected), then the entity parts are resolved and the period parsed. A failure
   there is one **row-level** reject and none of the row's measures load. A row with every cell
   empty is passed over silently.
4. Each measure cell: empty, or a placeholder from the measure's `nulls` (counted per
   placeholder in `null_skipped`), + `empty: skip` is skipped; otherwise ops then `scale` -> a
   fact, or a **cell-level** reject.
5. The same (entity, field, period) twice in one document (file + sheet): an equal value is
   dropped (the first cell is kept); a different value loads too, and the report names every
   cell involved. Across documents nothing is dropped: `add_facts` decides agreement and
   conflict at ingest, as it does between loads.

Before any sheet is read, every unit's `document_key` is rendered; two units with the same key
in one load are a `SpecMismatch` (one would supersede the other within the load). A
`SpecMismatch` on any unit refuses the whole load; with `source.sheets` it names the sheet.

After every unit, each open entity part's new names are compared pairwise for possible
duplicates (`resolve.similar_names`): `similar_entities` in the report and `similar.csv`.
Warnings only; they change no fact, count or gate.

**Native XLSX cells** (`cell_pipeline`): when a pipeline's first op is `number` and the cell is
a number, the cell's value is taken as is (`Decimal(raw)`) and only the later ops run, so a
`thousands: "."` spec doesn't misread the native 1234.5 as text. Likewise a first `date` op on
a date cell takes the cell's date, and a first `period` op on a date cell expands that date
with the op's `grain` and `fy_start_month`. Text cells (every CSV cell) go through the ops.

**Open entities and case.** The store resolves entity names case-insensitively, so within one
load an `open` value whose casefold matches an open value seen earlier, or a crosswalk
canonical, takes that first spelling (`Widget-a` then `WIDGET-A` -> both `Widget-a`); the
report counts it under `entities.normalised`. Composite names are compared casefolded too.

Every fact's `source_ref` is `<Sheet>!<A1>@wrangle:<spec>/<bundle_sha12>/<file_sha12>`, and
every sheet becomes one `RawDocument` whose text is a `# wrangle spec=<name> bundle=<sha256>`
line, then the sheet as CSV with Excel row numbers. The first line makes a corrected spec over
the same bytes a new document version, so `on_new_version="replace"` retires the old facts.
Values are `int` or, when not integral, `Decimal` (exact); `write_outputs` writes a Decimal as
a JSON number when a float reads back as the same value, else as a string.
"""

from __future__ import annotations

import csv
import io
import json
import re
from collections import Counter
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from ..rag.adapters.base import RawDocument, RawEntity, RawFact
from ..rag.store import FactRow
from .labels import header_fingerprint, header_labels
from .parse import ParseError, apply_text, compile_ops, period_of, to_fact_value
from .read import XLSX_SUFFIXES, Cell, Grid, col_letter, parse_ref, read_grid
from .resolve import STRATEGIES, FuzzyMatch, Resolver, similar_names
from .spec import PERIOD_JOIN, LoadedSpec, Measure, SkipBlank, SpecError, load_project_spec

__all__ = [
    "ApplyResult",
    "FactRecord",
    "FileResult",
    "SpecMismatch",
    "cell_pipeline",
    "document_key",
    "gate",
    "parse_source_ref",
    "reject_rate",
    "run",
    "unit_sheets",
    "write_outputs",
]

PREVIEW_FACTS = 50
_SOURCE_REF = re.compile(
    r"^(?P<sheet>.+)!(?P<cell>[A-Z]{1,3}[1-9][0-9]*)"
    r"@wrangle:(?P<spec>[a-z0-9_]+)/(?P<bundle>[0-9a-f]{12})/(?P<file>[0-9a-f]{12})$"
)
_READ_FAILURES = {"error": "error cell", "formula_no_value": "formula without cached value"}


class SpecMismatch(ValueError):
    """The file doesn't fit the spec (fingerprint, a missing column or sheet): the whole file is
    refused, rather than turned into rejects. With `source.sheets` the message names the sheet
    too, and one sheet's mismatch refuses the whole load."""

    def __init__(self, path: str, message: str, sheet: str | None = None):
        where = path if sheet is None else f"{path} sheet {sheet!r}"
        super().__init__(f"{where}: {message}")
        self.path = path
        self.sheet = sheet


def parse_source_ref(ref: str) -> dict[str, str]:
    """`Plan!G42@wrangle:planner_forecast/<bundle12>/<file12>` -> {sheet, cell, spec, bundle,
    file}. ValueError for anything else."""
    m = _SOURCE_REF.match(ref)
    if not m:
        raise ValueError(f"not a wrangle source_ref: {ref!r}")
    return m.groupdict()


# --- results --------------------------------------------------------------------------------


@dataclass(frozen=True)
class FactRecord:
    """A fact with where it came from: the cell (also in `source_ref`) plus the cells that gave
    its entity and period, which the report and preview show but the ref doesn't carry."""

    fact: RawFact
    file: str
    cell: str
    """"Sheet!A1"."""
    entity_cells: tuple[str, ...]
    period_from: str
    """The period's cell ("B7"), its cells joined with "+" when `period.column.name` lists
    several ("A6+B6"), or for an unpivoted cell its column label."""
    period_cells: tuple[str, ...] = ()
    """Every cell the period was read from ("A6", "B6"); empty for an unpivoted cell, whose
    period comes from its column label."""


@dataclass
class FileResult:
    path: str
    """As passed to `run`."""
    sheet: str
    file_sha256: str
    document: RawDocument
    facts: list[FactRecord] = field(default_factory=list)
    entities: list[RawEntity] = field(default_factory=list)
    """Composite and open entities with at least one fact in this file."""
    rows_read: int = 0
    rows_skipped: dict[str, int] = field(default_factory=dict)
    empty_skipped: int = 0
    null_skipped: dict[str, int] = field(default_factory=dict)
    """Measure cells skipped as a `nulls` placeholder, by placeholder."""
    section: dict[str, int] | None = None
    """With `start_after` / `stop_before`: the rows that matched them."""


@dataclass
class ApplyResult:
    loaded: LoadedSpec
    files: list[FileResult]
    crosswalk_entities: list[RawEntity]
    rejects: list[dict[str, Any]]
    report: dict[str, Any]
    proposed: list[FuzzyMatch]
    duplicates_dropped: int = 0
    conflicts: list[dict[str, Any]] = field(default_factory=list)
    similar_entities: list[dict[str, Any]] = field(default_factory=list)
    """Possible duplicate entity names (warnings only; see `_similar_entities`)."""

    @property
    def approved(self) -> bool:
        return self.loaded.approved

    @property
    def facts(self) -> list[RawFact]:
        return [r.fact for f in self.files for r in f.facts]

    @property
    def documents(self) -> list[RawDocument]:
        return [f.document for f in self.files]

    @property
    def entities(self) -> list[RawEntity]:
        """Crosswalk entities first, then each composite/open entity once."""
        out = list(self.crosswalk_entities)
        seen = {(e.kind, e.canonical_name) for e in out}
        for f in self.files:
            for e in f.entities:
                if (e.kind, e.canonical_name) not in seen:
                    seen.add((e.kind, e.canonical_name))
                    out.append(e)
        return out

    @property
    def reject_rate(self) -> float:
        return reject_rate(len(self.facts), len(self.rejects))


def reject_rate(facts: int, rejects: int) -> float:
    """rejects / (facts + rejects); 0 when nothing was read. Rule and empty skips don't count."""
    total = facts + rejects
    return rejects / total if total else 0.0


# --- compiling the spec ---------------------------------------------------------------------

Pipeline = Callable[[Cell], object]


def _cell_date(cell: Cell, op: str) -> date:
    try:  # `2025-10-31`, or `2025-10-31T13:00:00` for a date cell with a time of day
        return date.fromisoformat(cell.raw[:10])
    except ValueError as e:
        raise ParseError(f"{op}: unparseable", op) from e


def cell_pipeline(ops: list[dict[str, Any]]) -> Pipeline:
    """`compile_ops` for a cell rather than its raw string: the native-cell rule in the module
    docstring. ValueError for a bad pipeline, as `compile_ops`."""
    whole = compile_ops(ops)
    # Leading `text` ops only trim and case-fold; on a native number or date cell they are
    # no-ops, so they don't stop the native rule (`[text, number]` is a common spec).
    k = 0
    while k < len(ops) - 1 and ops[k]["op"] == "text":
        k += 1
    first = ops[k]
    rest = compile_ops(ops[k + 1 :]) if len(ops) > k + 1 else None
    name = first["op"]

    def run(cell: Cell) -> object:
        value: object
        if name == "number" and cell.kind == "number":
            try:
                value = Decimal(cell.raw)
            except InvalidOperation as e:
                raise ParseError("number: unparseable", "number") from e
        elif name == "date" and cell.kind == "date":
            value = _cell_date(cell, "date")
        elif name == "period" and cell.kind == "date":
            value = period_of(
                _cell_date(cell, "period"), first["grain"], first.get("fy_start_month", 1)
            )
        else:
            return whole(cell.raw)
        return rest(value) if rest is not None else value  # type: ignore[arg-type]

    return run


def _compile(ops: list[dict[str, Any]], where: str) -> Pipeline | None:
    if not ops:
        return None
    try:
        return cell_pipeline(ops)
    except ValueError as e:
        raise SpecError(f"{where}: {e}") from e


def _period_pipeline(ops: list[dict[str, Any]], where: str, grain: str | None) -> Pipeline:
    if not ops:
        raise SpecError(f"{where}: needs ops ending in `date` or `period`")
    last = ops[-1]["op"]
    if last not in ("date", "period"):
        raise SpecError(f"{where}: ops must end in `date` or `period`, not {last!r}")
    if last == "period" and grain is not None:
        raise SpecError(f"{where}: `grain` goes in the `period` op, not beside it")
    fn = _compile(ops, where)
    assert fn is not None

    def run(cell: Cell) -> tuple[date, date]:
        v = fn(cell)
        if isinstance(v, tuple):
            return v
        assert isinstance(v, date)
        return period_of(v, grain or "day")

    return run


@dataclass
class _Compiled:
    parts: list[tuple[Pipeline | None, Resolver]]
    period_column: Pipeline | None
    period_cell: Pipeline | None
    unpivot: Pipeline | None
    measures: list[tuple[Measure, Pipeline | None]]
    skips: list[tuple[str, Any]]


def _compile_spec(loaded: LoadedSpec, names: dict[str, str]) -> _Compiled:
    """`names` (casefold -> first spelling) is shared by every part's resolver, for the
    case rule on open values."""
    spec = loaded.spec
    parts = [
        (
            _compile(p.ops, f"entity.parts[{i}]"),
            Resolver.for_part(p.resolve, loaded.crosswalks, names=names),
        )
        for i, p in enumerate(spec.entity.parts)
    ]
    per = spec.period
    measures = []
    for i, m in enumerate(spec.measures):
        if not m.is_text and m.ops[-1]["op"] != "number":
            raise SpecError(f"measures[{i}] ({m.field}): numeric measures end in `number`")
        measures.append((m, _compile(m.ops, f"measures[{i}]")))
    skips: list[tuple[str, Any]] = []
    for rule in spec.source.skip_rows:
        if isinstance(rule, SkipBlank):
            skips.append(("blank", rule))
        else:
            skips.append(("match", (rule, re.compile(rule.matches))))
    return _Compiled(
        parts=parts,
        period_column=(
            _period_pipeline(per.column.ops, "period.column", per.column.grain)
            if per.column
            else None
        ),
        period_cell=_period_pipeline(per.cell.ops, "period.cell", None) if per.cell else None,
        unpivot=_period_pipeline(per.unpivot.ops, "period.unpivot", None) if per.unpivot else None,
        measures=measures,
        skips=skips,
    )


# --- per file -------------------------------------------------------------------------------


class _Columns:
    """Joined header labels -> 1-based column numbers."""

    def __init__(self, path: str, labels: list[str], sheet: str | None = None):
        self.path = path
        self.sheet = sheet
        self.labels = labels
        self._exact: dict[str, list[int]] = {}
        self._loose: dict[str, list[int]] = {}
        for i, label in enumerate(labels, start=1):
            self._exact.setdefault(label, []).append(i)
            self._loose.setdefault(apply_text(label), []).append(i)

    def find(self, name: str, what: str) -> int:
        cols = self._exact.get(name) or self._loose.get(apply_text(name), [])
        if not cols:
            shown = [lbl for lbl in self.labels if lbl]
            raise SpecMismatch(
                self.path, f"{what} {name!r} is not a column; labels: {shown}", self.sheet
            )
        if len(cols) > 1:
            letters = [col_letter(c) for c in cols]
            raise SpecMismatch(
                self.path, f"{what} {name!r} names several columns: {letters}", self.sheet
            )
        return cols[0]


def _render_document(grid: Grid, spec: str, bundle_sha256: str) -> str:
    buf = io.StringIO()
    buf.write(f"# wrangle spec={spec} bundle={bundle_sha256}\n")
    w = csv.writer(buf, lineterminator="\n")
    width = len(grid.rows[0]) if grid.rows else 0
    w.writerow(["row", *(col_letter(c) for c in range(1, width + 1))])
    for i, row in enumerate(grid.rows, start=1):
        w.writerow([i, *(c.raw for c in row)])
    return buf.getvalue()


def _period_label(start: date, end: date) -> str:
    if start == end:
        return start.isoformat()
    if start.day == 1 and period_of(start, "month") == (start, end):
        return start.strftime("%Y-%m")
    return f"{start.isoformat()}..{end.isoformat()}"


def _precision(m: Measure) -> int | float | None:
    if m.is_text:
        return None
    p = m.effective_precision
    if p is not None and float(p).is_integer():
        return int(p)
    return p


def _is_xlsx(path: Path | str) -> bool:
    return Path(path).suffix.lower() in XLSX_SUFFIXES


def _workbook_sheets(path: Path | str) -> list[str]:
    """The workbook's worksheet names in workbook order (chartsheets left out). Only the
    workbook part is parsed, no sheet."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # zipfile.BadZipFile, KeyError, InvalidFileException, OSError
        raise SpecMismatch(str(path), f"cannot open the workbook: {e}") from e
    try:
        return [ws.title for ws in wb.worksheets]
    finally:
        wb.close()


def unit_sheets(loaded: LoadedSpec, path: Path | str) -> list[str | None]:
    """The sheets the spec reads from one file, in workbook order: `[source.sheet]`, the
    `source.sheets` selection (which opens the workbook to list its sheets), or `[None]` for a
    CSV. `SpecMismatch` for an XLSX with neither, a CSV with either, or a selection that names
    a missing sheet or matches none."""
    src = loaded.spec.source
    if not _is_xlsx(path):
        if src.sheet is not None or src.sheets is not None:
            which = "source.sheet" if src.sheet is not None else "source.sheets"
            raise SpecMismatch(
                str(path), f"is a CSV, whose only sheet is its file stem: omit {which}"
            )
        return [None]
    if src.sheets is not None:
        try:
            return list(src.select_sheets(_workbook_sheets(path)))
        except ValueError as e:
            if isinstance(e, SpecMismatch):
                raise
            raise SpecMismatch(str(path), str(e)) from e
    if src.sheet is None:
        raise SpecMismatch(str(path), "XLSX needs source.sheet or source.sheets in the spec")
    return [src.sheet]


def document_key(loaded: LoadedSpec, path: Path | str, sheet: str | None = None) -> str:
    """The spec's `document_key` rendered for one file (and, with `source.sheets`, one of its
    sheets, which must then be given). Known before the sheet is read: the sheet is `sheet`,
    else `source.sheet` for XLSX (`SpecMismatch` when the spec has neither) and the file stem
    for CSV."""
    src = loaded.spec.source
    if sheet is None:
        if _is_xlsx(path):
            if src.sheets is not None:
                raise SpecMismatch(
                    str(path), "the spec reads several sheets (source.sheets): name one"
                )
            if src.sheet is None:
                raise SpecMismatch(
                    str(path), "XLSX needs source.sheet or source.sheets in the spec"
                )
            sheet = src.sheet
        else:
            sheet = Path(path).stem
    return src.document_key.format(
        spec=loaded.spec.spec, file=Path(path).name, path=str(path), sheet=sheet
    )


def _others(canonicals: list[str], i: int) -> tuple[str, ...]:
    """The scope a part's names are compared within: the parts to its left, read as a
    hierarchy (state, then county). Parts to the right don't scope it: in plant / part /
    line, two spellings of one part number under the same plant are still worth a look."""
    return tuple(canonicals[:i])


class _Run:
    def __init__(self, loaded: LoadedSpec, force: bool):
        self.loaded = loaded
        self.spec = loaded.spec
        self.force = force
        canonicals = sorted(c for cw in loaded.crosswalks.values() for c in cw.canonicals)
        # casefold -> the first spelling this load used (crosswalk canonicals first).
        self.names: dict[str, str] = {}
        for c in canonicals:
            self.names.setdefault(c.casefold(), c)
        self.crosswalk_canonicals = {c.casefold() for c in canonicals}
        self.c = _compile_spec(loaded, self.names)
        self.rejects: list[dict[str, Any]] = []
        self.duplicates = 0
        # Per document: (entity, field, start, end) -> [(value, "Sheet!A1")]; values within a
        # key all differ. Reset for each file; `conflicts` collects across them.
        self.seen: dict[tuple[str, str, date, date], list[tuple[Any, str]]] = {}
        self.conflicts: list[dict[str, Any]] = []
        # Per entity part: is it open, and each new (non-crosswalk) canonical with at least
        # one fact -> the cell it was first read from, keyed by the parts to its left. In a
        # composite entity only names under the same left parts are compared: `Johnson
        # County` in Kansas is not a near-duplicate of `Johnston County` in North Carolina.
        self.open_parts = [p.resolve.policy == "open" for p in self.spec.entity.parts]
        self.part_names: list[dict[tuple[tuple[str, ...], str], str]] = [
            {} for _ in self.spec.entity.parts
        ]

    def similar_entities(self) -> list[dict[str, Any]]:
        """Possible duplicates among each open part's new canonicals, compared per part (never
        the composite string): `resolve.similar_warn` and `not_same` (§7.3). Warnings only."""
        out: list[dict[str, Any]] = []
        kind = self.spec.entity.kind
        for i, part in enumerate(self.spec.entity.parts):
            if not self.open_parts[i]:
                continue
            groups: dict[tuple[str, ...], dict[str, str]] = {}
            for (ctx, name), where in self.part_names[i].items():
                groups.setdefault(ctx, {})[name] = where
            for ctx, names in groups.items():
                for p in similar_names(
                    list(names), part.resolve.similar_warn, not_same=part.resolve.not_same
                ):
                    item = {
                        "part": part.column,
                        "kind": kind,
                        "a": p.a,
                        "b": p.b,
                        "score": p.score,
                        "reason": p.reason,
                        "cells": [names[p.a], names[p.b]],
                    }
                    if ctx:
                        item["within"] = list(ctx)
                    out.append(item)
        return sorted(out, key=lambda d: (-d["score"], d["a"], d["b"], d["part"]))

    def _reject(self, fr: FileResult, *, where: dict[str, Any], **kw: Any) -> None:
        rec: dict[str, Any] = {"file": fr.path, "sheet": fr.sheet, **where}
        rec.update(kw)
        self.rejects.append(rec)

    def _row_reject(self, fr, row: int, cell: Cell, stage: str, op: str, reason: str) -> None:
        self._reject(
            fr,
            where={"row": row, "level": "row"},
            raw=cell.raw,
            stage=stage,
            op=op,
            reason=reason,
        )

    def _cell_reject(self, fr, cell: Cell, raw: str, stage: str, op: str, reason: str) -> None:
        self._reject(
            fr,
            where={"cell": cell.ref, "level": "cell"},
            raw=raw,
            stage=stage,
            op=op,
            reason=reason,
        )

    def file(self, path: Path | str, sheet: str | None = None) -> FileResult:
        """One unit: a CSV, an XLSX file's `source.sheet`, or (with `source.sheets`) one of the
        selected sheets, passed as `sheet`."""
        spec, src = self.spec, self.spec.source
        shown = str(path)
        if sheet is None and _is_xlsx(path):
            sheet = src.sheet
        source = document_key(self.loaded, path, sheet)
        # Mismatch messages name the sheet when the spec reads several.
        msheet = sheet if src.sheets is not None else None
        try:
            grid = read_grid(
                path,
                sheet if _is_xlsx(path) else None,
                encoding=src.encoding,
                delimiter=src.delimiter,
            )
        except ValueError as e:
            raise SpecMismatch(shown, str(e), msheet) from e

        labels = header_labels(grid, list(src.header_rows))
        if src.header_fingerprint and not self.force:
            got = header_fingerprint(labels)
            if got != src.header_fingerprint:
                raise SpecMismatch(
                    shown,
                    f"header_fingerprint is {got}, the spec expects {src.header_fingerprint} "
                    "(the header changed; re-check the spec, or --force)",
                    msheet,
                )
        cols = _Columns(shown, labels, msheet)
        part_cols = [cols.find(p.column, "entity column") for p in spec.entity.parts]
        skip_cols: list[tuple[str, Any]] = []
        for kind, rule in self.c.skips:
            if kind == "blank":
                skip_cols.append(
                    (kind, [cols.find(n, "skip_rows blank column") for n in rule.blank])
                )
            else:
                r, rx = rule
                skip_cols.append((kind, (cols.find(r.column, "skip_rows column"), rx)))
        per = spec.period
        period_cols = (
            [cols.find(n, "period column") for n in per.column.names] if per.column else []
        )
        first_row, last_row = self._section(grid, cols, shown, msheet)
        unpivot_cols: list[int] = []
        if per.unpivot is not None:
            if per.unpivot.columns is not None:
                lo = cols.find(per.unpivot.columns.from_, "unpivot.columns.from")
                hi = cols.find(per.unpivot.columns.to, "unpivot.columns.to")
                if lo > hi:
                    raise SpecMismatch(shown, "unpivot.columns.from is right of `to`", msheet)
                unpivot_cols = list(range(lo, hi + 1))
            else:
                rx = re.compile(per.unpivot.matching or "")
                unpivot_cols = [i for i, lbl in enumerate(labels, 1) if rx.search(lbl)]
                if not unpivot_cols:
                    raise SpecMismatch(
                        shown, f"no column label matches {per.unpivot.matching!r}", msheet
                    )
        measure_cols = [
            (m, fn, cols.find(m.column, "measure column") if m.column is not None else None)
            for m, fn in self.c.measures
        ]

        document = RawDocument(
            source=source,
            title=f"{Path(path).name} — {grid.sheet}",
            text=_render_document(grid, spec.spec, self.loaded.bundle_sha256),
            metadata={
                "file_sha256": grid.file_sha256,
                "spec": spec.spec,
                "bundle_sha256": self.loaded.bundle_sha256,
                "header_rows": json.dumps(list(src.header_rows)),
            },
        )
        fr = FileResult(
            path=shown,
            sheet=grid.sheet,
            file_sha256=grid.file_sha256,
            document=document,
            rows_skipped={f"skip_rows[{i}]": 0 for i in range(len(self.c.skips))},
        )
        if src.start_after is not None or src.stop_before is not None:
            fr.section = {}
            if src.start_after is not None:
                fr.section["start_after_row"] = first_row - 1
            if src.stop_before is not None:
                fr.section["stop_before_row"] = last_row + 1
        ref_suffix = f"@wrangle:{spec.spec}/{self.loaded.bundle_sha12}/{grid.file_sha256[:12]}"

        # Periods that don't depend on the row: a fixed cell, or each unpivoted column's label.
        fixed: tuple[Cell, tuple[date, date] | None, ParseError | None] | None = None
        if per.cell is not None:
            assert self.c.period_cell is not None
            r, c = parse_ref(per.cell.ref)
            fixed = (grid.cell(r, c), *self._period(grid.cell(r, c), self.c.period_cell))
        label_periods: dict[int, tuple[date, date] | ParseError] = {}
        for col in unpivot_cols:
            assert self.c.unpivot is not None
            try:  # a label is text, whatever the header cells were
                label_periods[col] = self.c.unpivot(Cell(0, col, labels[col - 1], "text"))
            except ParseError as e:
                label_periods[col] = e

        entities: dict[str, RawEntity] = {}
        self.seen = {}
        for r in range(first_row, last_row + 1):
            row = grid.rows[r - 1]
            if all(c.kind == "empty" for c in row):
                continue
            fr.rows_read += 1
            if self._skipped(grid, r, skip_cols, fr):
                continue

            # Entity parts -> canonical.
            canonicals: list[str] = []
            failed = False
            for (fn, resolver), col in zip(self.c.parts, part_cols, strict=True):
                cell = grid.cell(r, col)
                if cell.kind in _READ_FAILURES:
                    self._row_reject(fr, r, cell, "entity", "read", _READ_FAILURES[cell.kind])
                    failed = True
                    break
                try:
                    cleaned = str(fn(cell)) if fn else cell.raw
                except ParseError as e:
                    self._row_reject(fr, r, cell, "entity", e.op or "", e.reason)
                    failed = True
                    break
                res = resolver.resolve(cleaned)
                if res.canonical is None:
                    self._row_reject(fr, r, cell, "entity", "resolve", res.reason or "")
                    failed = True
                    break
                canonicals.append(res.canonical)
            if failed:
                continue
            # Each open part's new name and its cell, noted once the row loads a fact.
            new_names = [
                (i, (_others(canonicals, i), canonical), f"{grid.sheet}!{col_letter(col)}{r}")
                for i, (canonical, col) in enumerate(zip(canonicals, part_cols, strict=True))
                if self.open_parts[i]
                and (_others(canonicals, i), canonical) not in self.part_names[i]
                and canonical.casefold() not in self.crosswalk_canonicals
            ]
            entity = spec.entity.render(canonicals)
            entity = self.names.setdefault(entity.casefold(), entity)
            entity_cells = tuple(f"{col_letter(c)}{r}" for c in part_cols)

            # Period for the row (column or fixed cell); unpivot is per cell, below.
            row_period: tuple[date, date] | None = None
            err: ParseError | None = None
            period_cells: tuple[str, ...] = ()
            if period_cols:
                assert self.c.period_column is not None
                pcell = self._period_cell(grid, r, period_cols)
                row_period, err = self._period(pcell, self.c.period_column)
                period_cells = tuple(f"{col_letter(c)}{r}" for c in period_cols)
            elif fixed is not None:
                pcell, row_period, err = fixed
                period_cells = (pcell.ref,)
            period_from = "+".join(period_cells)
            if err is not None:
                self._row_reject(fr, r, pcell, "period", err.op or "", err.reason)
                continue

            for m, fn, mcol in measure_cols:
                targets = [mcol] if mcol is not None else unpivot_cols
                for col in targets:
                    cell = grid.cell(r, col)
                    null = self._null(cell, m)
                    if (cell.kind == "empty" or null is not None) and m.empty == "skip":
                        if null is None:
                            fr.empty_skipped += 1
                        else:
                            fr.null_skipped[null] = fr.null_skipped.get(null, 0) + 1
                        continue
                    pcells = period_cells
                    if row_period is not None:
                        period, pfrom = row_period, period_from
                    else:
                        pcells = ()
                        lp = label_periods[col]
                        if isinstance(lp, ParseError):
                            self._cell_reject(
                                fr, cell, labels[col - 1], "period", lp.op or "", lp.reason
                            )
                            continue
                        period, pfrom = lp, labels[col - 1]
                    if null is not None:  # empty: reject
                        self._cell_reject(
                            fr, cell, cell.raw, "measure", "nulls", "null placeholder"
                        )
                        continue
                    value = self._measure_value(fr, cell, m, fn)
                    if value is None:
                        continue
                    cell_name = f"{grid.sheet}!{cell.ref}"
                    if not self._keep(entity, m.field, period, value, cell_name):
                        continue
                    row_ = FactRow(
                        entity=entity,
                        field=m.field,
                        value=value,
                        unit=self.loaded.unit_for(m),
                        period_start=period[0],
                        period_end=period[1],
                        precision=_precision(m),  # type: ignore[arg-type]
                        source_ref=cell_name + ref_suffix,
                    )
                    fr.facts.append(
                        FactRecord(
                            RawFact(source, row_), shown, cell_name, entity_cells, pfrom, pcells
                        )
                    )
                    if (
                        entity.casefold() not in self.crosswalk_canonicals
                        and entity not in entities
                    ):
                        entities[entity] = RawEntity(spec.entity.kind, entity)
                    for i, key, where in new_names:
                        self.part_names[i].setdefault(key, where)
                    new_names = []
        fr.entities = list(entities.values())
        self.conflicts.extend(
            {
                "document": source,
                "key": [k[0], k[1], _period_label(k[2], k[3])],
                "cells": [c for _, c in vs],
            }
            for k, vs in self.seen.items()
            if len(vs) > 1
        )
        return fr

    def _section(
        self, grid: Grid, cols: _Columns, shown: str, msheet: str | None
    ) -> tuple[int, int]:
        """(first, last) data row, inclusive: `first_data_row` to the grid's end, narrowed by
        `start_after` and `stop_before`. A bound that matches no row is a `SpecMismatch`, since
        loading on regardless would read the wrong table (or both)."""
        src = self.spec.source
        first, last = src.first_data_row, len(grid.rows)
        if src.start_after is not None:
            hit = self._first_match(grid, cols, src.start_after, first, last, "start_after")
            if hit is None:
                raise SpecMismatch(
                    shown,
                    f"source.start_after: no row from {first} has {src.start_after.column!r} "
                    f"matching {src.start_after.matches!r}",
                    msheet,
                )
            first = hit + 1
        if src.stop_before is not None:
            hit = self._first_match(grid, cols, src.stop_before, first, last, "stop_before")
            if hit is None:
                raise SpecMismatch(
                    shown,
                    f"source.stop_before: no row from {first} has {src.stop_before.column!r} "
                    f"matching {src.stop_before.matches!r}",
                    msheet,
                )
            last = hit - 1
        return first, last

    @staticmethod
    def _first_match(grid: Grid, cols: _Columns, bound, lo: int, hi: int, what: str) -> int | None:
        col = cols.find(bound.column, f"source.{what} column")
        rx = re.compile(bound.matches)
        for r in range(lo, hi + 1):
            if rx.search(grid.cell(r, col).raw) is not None:
                return r
        return None

    @staticmethod
    def _period_cell(grid: Grid, r: int, period_cols: list[int]) -> Cell:
        """The period column's cell; with several columns, a text cell (at the first one) whose
        raw is their whitespace-cleaned values joined with `PERIOD_JOIN`. An error or
        formula-without-value cell among them is returned as is, so it rejects as a read."""
        cells = [grid.cell(r, c) for c in period_cols]
        if len(cells) == 1:
            return cells[0]
        for c in cells:
            if c.kind in _READ_FAILURES:
                return c
        if all(c.kind == "empty" for c in cells):
            return Cell(r, period_cols[0], "", "empty")
        return Cell(r, period_cols[0], PERIOD_JOIN.join(apply_text(c.raw) for c in cells), "text")

    @staticmethod
    def _null(cell: Cell, m: Measure) -> str | None:
        """The `nulls` placeholder this cell holds, or None."""
        if not m.nulls or cell.kind in _READ_FAILURES or cell.kind == "empty":
            return None
        s = cell.raw.strip()
        return s if s in m.nulls else None

    def _skipped(self, grid: Grid, r: int, skip_cols, fr: FileResult) -> bool:
        for i, (kind, how) in enumerate(skip_cols):
            if kind == "blank":
                hit = all(grid.cell(r, c).kind == "empty" for c in how)
            else:
                col, rx = how
                hit = rx.search(grid.cell(r, col).raw) is not None
            if hit:
                fr.rows_skipped[f"skip_rows[{i}]"] += 1
                return True
        return False

    @staticmethod
    def _period(cell: Cell, fn: Pipeline) -> tuple[tuple[date, date] | None, ParseError | None]:
        """(period, None), or (None, the failure)."""
        if cell.kind in _READ_FAILURES:
            return None, ParseError(_READ_FAILURES[cell.kind], "read")
        try:
            return fn(cell), None  # type: ignore[return-value]
        except ParseError as e:
            return None, e

    def _measure_value(
        self, fr: FileResult, cell: Cell, m: Measure, fn: Pipeline | None
    ) -> int | Decimal | str | None:
        if cell.kind in _READ_FAILURES:
            self._cell_reject(fr, cell, cell.raw, "measure", "read", _READ_FAILURES[cell.kind])
            return None
        try:
            out = fn(cell) if fn else cell.raw
        except ParseError as e:
            self._cell_reject(fr, cell, cell.raw, "measure", e.op or "", e.reason)
            return None
        if cell.kind == "empty":  # empty: reject, and the ops didn't already refuse it
            self._cell_reject(fr, cell, cell.raw, "measure", "read", "empty cell")
            return None
        if m.is_text:
            return str(out)
        assert isinstance(out, Decimal)
        return to_fact_value(out * m.scale_decimal)

    def _keep(self, entity: str, fld: str, period, value, cell_name: str) -> bool:
        key = (entity, fld, period[0], period[1])
        prior = self.seen.setdefault(key, [])
        if any(v == value for v, _ in prior):
            self.duplicates += 1
            return False
        prior.append((value, cell_name))
        return True


def run(
    spec: LoadedSpec | str | Path, files: Iterable[str | Path], *, force: bool = False
) -> ApplyResult:
    """The dry run: every file through the spec, nothing written anywhere.

    Raises `SpecMismatch` when a file doesn't fit the spec (header fingerprint, unless `force`;
    a missing column or sheet) and `SpecError` for a spec whose ops can't run."""
    loaded = spec if isinstance(spec, LoadedSpec) else load_project_spec(spec)
    files = list(files)
    units: list[tuple[str | Path, str | None]] = [
        (f, sheet) for f in files for sheet in unit_sheets(loaded, f)
    ]
    keys: dict[str, str] = {}
    for f, sheet in units:
        key = document_key(loaded, f, sheet)
        shown = (
            str(f) if sheet is None or loaded.spec.source.sheets is None else f"{f} sheet {sheet!r}"
        )
        if key in keys:
            raise SpecMismatch(
                str(f),
                f"document_key {loaded.spec.source.document_key!r} renders {key!r} for this "
                f"{'sheet' if loaded.spec.source.sheets is not None else 'file'} and for "
                f"{keys[key]}: in one load they would supersede each other "
                "(load them separately, or use a document_key that tells them apart)",
                sheet if loaded.spec.source.sheets is not None else None,
            )
        keys[key] = shown
    state = _Run(loaded, force)
    results = [state.file(f, sheet) for f, sheet in units]
    similar = state.similar_entities()

    crosswalk_entities = [
        RawEntity(cw.kind, canonical, cw.aliases_for(canonical))
        for cw in loaded.crosswalks.values()
        for canonical in sorted(cw.canonicals)
    ]
    counts = dict.fromkeys(STRATEGIES, 0)
    proposed: list[FuzzyMatch] = []
    for _, resolver in state.c.parts:
        for k, v in resolver.counts.items():
            counts[k] += v
        proposed.extend(resolver.fuzzy_matches)

    n_facts = sum(len(f.facts) for f in results)
    n_rejects = len(state.rejects)
    conflicts = state.conflicts
    report = {
        "spec": loaded.spec.spec,
        "bundle_sha256": loaded.bundle_sha256,
        "approved": loaded.approved,
        "files": [
            {
                "path": f.path,
                "sheet": f.sheet,
                "file_sha256": f.file_sha256,
                "rows_read": f.rows_read,
                "rows_skipped": dict(f.rows_skipped),
                "empty_skipped": f.empty_skipped,
                "null_skipped": dict(sorted(f.null_skipped.items())),
                **({"section": dict(f.section)} if f.section is not None else {}),
            }
            for f in results
        ],
        "facts": n_facts,
        "duplicates_dropped": state.duplicates,
        "rejects": {
            "total": n_rejects,
            "rate": round(reject_rate(n_facts, n_rejects), 6),
            "by_reason": dict(sorted(Counter(r["reason"] for r in state.rejects).items())),
        },
        "entities": counts,
        "conflicts_within_load": conflicts,
        "similar_entities": similar,
    }
    return ApplyResult(
        loaded=loaded,
        files=results,
        crosswalk_entities=crosswalk_entities,
        rejects=state.rejects,
        report=report,
        proposed=proposed,
        duplicates_dropped=state.duplicates,
        conflicts=conflicts,
        similar_entities=similar,
    )


# --- gates and outputs ----------------------------------------------------------------------


def gate(
    result: ApplyResult,
    *,
    approved: bool | None = None,
    draft: bool = False,
    max_reject_rate: float = 0.02,
) -> list[str]:
    """Why apply must refuse (§8), or [] to go ahead. `approved` defaults to the spec's own
    state; `draft` waives the approval check only. (A `SpecMismatch` is raised by `run`.)"""
    reasons = []
    is_approved = result.approved if approved is None else approved
    if not draft and not is_approved:
        approval = result.loaded.spec.approval
        if approval is None:
            reasons.append(
                f"spec {result.loaded.spec.spec!r} is not approved "
                "(run `fde wrangle approve`, or load with --draft)"
            )
        else:
            reasons.append(
                f"spec {result.loaded.spec.spec!r} approval is stale: approved at bundle "
                f"{approval.bundle_sha256[:12]}, the bundle is now {result.loaded.bundle_sha12}"
            )
    rate = result.reject_rate
    if rate > max_reject_rate:
        reasons.append(
            f"reject rate {rate:.4f} ({len(result.rejects)} rejects) is over the limit "
            f"{max_reject_rate}"
        )
    return reasons


def _json_default(o: object) -> object:
    """A Decimal value: a JSON number when the float reads back as the same value, else a
    string, so no digit is lost either way."""
    if isinstance(o, Decimal):
        f = float(o)
        return f if Decimal(repr(f)) == o else str(o)
    raise TypeError(f"{type(o).__name__} is not JSON serialisable")


def _dumps(obj: object, **kw: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, default=_json_default, **kw)


def _fact_json(rec: FactRecord) -> dict[str, Any]:
    row = rec.fact.row
    return {
        "entity": row.entity,
        "field": row.field,
        "value": row.value,
        "unit": row.unit,
        "period_start": row.period_start.isoformat() if row.period_start else None,
        "period_end": row.period_end.isoformat() if row.period_end else None,
        "precision": row.precision,
        "source_ref": row.source_ref,
        "document": rec.fact.document_source,
        "entity_cells": list(rec.entity_cells),
        "period_from": rec.period_from,
        "period_cells": list(rec.period_cells),
    }


def write_outputs(result: ApplyResult, out_dir: str | Path) -> dict[str, Path]:
    """report.json, rejects.jsonl, facts.preview.jsonl (first 50), crosswalk.proposed.csv and
    similar.csv (possible duplicate entities)."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths = {
        "report": out / "report.json",
        "rejects": out / "rejects.jsonl",
        "preview": out / "facts.preview.jsonl",
        "proposed": out / "crosswalk.proposed.csv",
        "similar": out / "similar.csv",
    }
    paths["report"].write_text(_dumps(result.report, indent=2) + "\n", encoding="utf-8")
    paths["rejects"].write_text(
        "".join(_dumps(r) + "\n" for r in result.rejects),
        encoding="utf-8",
    )
    records = [r for f in result.files for r in f.facts][:PREVIEW_FACTS]
    paths["preview"].write_text(
        "".join(_dumps(_fact_json(r)) + "\n" for r in records),
        encoding="utf-8",
    )
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["crosswalk", "alias", "canonical", "score", "runner_up"])
    for p in result.proposed:
        w.writerow([p.crosswalk, p.alias, p.canonical, f"{p.score:g}", f"{p.second:g}"])
    paths["proposed"].write_text(buf.getvalue(), encoding="utf-8")
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["part", "kind", "within", "a", "b", "score", "reason", "cell_a", "cell_b"])
    for d in result.similar_entities:
        w.writerow(
            [
                d["part"],
                d["kind"],
                " / ".join(d.get("within", [])),
                d["a"],
                d["b"],
                f"{d['score']:g}",
                d["reason"],
                *d["cells"],
            ]
        )
    paths["similar"].write_text(buf.getvalue(), encoding="utf-8")
    return paths
